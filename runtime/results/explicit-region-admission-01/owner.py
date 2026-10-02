import hashlib,json,os,signal,socket,subprocess,sys,time
from pathlib import Path
from Xlib import display,X,error
from openpyxl import Workbook,load_workbook
archive,case=map(lambda p:Path(p).resolve(),sys.argv[1:3])
seed=int(sys.argv[3]);reuse=False
phase='feedback'
if len(sys.argv)!=6:raise ValueError('archive case seed compact wait arguments required')
if sys.argv[4]!='compact':raise ValueError('fixed public compact route')
wait_ms=int(sys.argv[5])
schedule=json.loads((Path(__file__).parent/'schedule.json').read_text())
rows=[row for row in schedule if row['seed']==seed]
if len(rows)!=1 or rows[0]['wait_ms']!=wait_ms or case.name!=rows[0]['case']:raise ValueError('frozen allocation mismatch')
expected=rows[0]['values'];max_commands=12
case.mkdir(exist_ok=False)
sys.path.insert(0,str(archive))
from runtime.guarded_x11_v1.bridge import read_window_title
from delivered_capture import DeliveredCapture
from capture_request import region_for
from runtime.guarded_x11_v1.handles import ALIAS
NATIVE_ALIASES=('sheet_context','format_context')
if any(ALIAS.fullmatch(alias) is None for alias in NATIVE_ALIASES):
    raise ValueError('invalid native aliases before display allocation')
def save(name,row):(case/name).write_text(json.dumps(row,indent=2)+'\n')
children=[];host_exit=None;connection=None;bridge=None;owner=None
save('allocation.json',{'seed':seed,'reuseReviewedImages':reuse,'pid':os.getpid(),'source':json.loads((archive.parent/'MANIFEST.json').read_text())['source_revision'],'archive_sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'retry_budget':0,'expected':expected,'route':sys.argv[4],'phase':phase,'repair_budget':0,'post_release_wait_ms':wait_ms})
for n in range(27000,27100):
    if Path(f'/tmp/.X11-unix/X{n}').exists() or Path(f'/tmp/.X{n}-lock').exists():continue
    sock=socket.socket(socket.AF_UNIX)
    try:sock.connect('\0'+f'/tmp/.X11-unix/X{n}')
    except OSError:break
    finally:sock.close()
else:raise RuntimeError('no unused display')
env={k:v for k,v in os.environ.items() if k not in ('WAYLAND_DISPLAY','DBUS_SESSION_BUS_ADDRESS')}
env.update(DISPLAY=f':{n}',GDK_BACKEND='x11',QT_QPA_PLATFORM='xcb',LANG='C.UTF-8',LC_ALL='C.UTF-8',SAL_USE_VCLPLUGIN='gen')
for key,folder in [('HOME','home'),('XDG_CONFIG_HOME','config'),('XDG_CACHE_HOME','cache'),('XDG_DATA_HOME','data'),('XDG_RUNTIME_DIR','run')]:
    path=case/folder;path.mkdir(mode=0o700);env[key]=str(path)
def launch(label,args,stdin=subprocess.DEVNULL):
    p=subprocess.Popen(args,env=env,stdin=stdin,stdout=(case/(label+'.stdout')).open('xb'),stderr=(case/(label+'.stderr')).open('xb'),start_new_session=True)
    children.append((label,p));return p
try:
    xvfb=launch('xvfb',['Xvfb',f':{n}','-screen','0','1280x800x24','-ac','-nolisten','tcp'])
    deadline=time.monotonic()+5
    while True:
        if xvfb.poll() is not None:raise RuntimeError('Xvfb exited')
        try:connection=display.Display(f':{n}');break
        except Exception:
            if time.monotonic()>deadline:raise
            time.sleep(.02)
    wm=launch('openbox',['openbox','--config-file','/etc/xdg/openbox/rc.xml'])
    atom=connection.intern_atom('_NET_SUPPORTING_WM_CHECK')
    deadline=time.monotonic()+5
    while connection.screen().root.get_full_property(atom,X.AnyPropertyType) is None:
        if wm.poll() is not None or time.monotonic()>deadline:raise RuntimeError('WM not ready')
        time.sleep(.02)
    workbook=case/f'sheet-{seed}.xlsx';Workbook().save(workbook)
    save('original-workbook.json',{'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    profile=case/'lo-profile';(profile/'user').mkdir(parents=True)
    (profile/'user'/'registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="ShowTipOfTheDay" oor:op="fuse"><value>false</value></prop></item></oor:items>')
    app=launch('calc',['libreoffice','--norestore','--nodefault','--nolockcheck',f'-env:UserInstallation={profile.as_uri()}','--calc',str(workbook)])
    deadline=time.monotonic()+30;target=None
    while target is None:
        prop=connection.screen().root.get_full_property(connection.intern_atom('_NET_CLIENT_LIST'),X.AnyPropertyType)
        for wid in ([] if prop is None else prop.value):
            win=connection.create_resource_object('window',int(wid))
            try:
                legacy=win.get_wm_name()
                title=read_window_title(connection,win) or ''
            except error.BadWindow as stale_error:
                with (case/'window-discovery-refusals.jsonl').open('a') as sample:
                    sample.write(json.dumps({'ns':time.monotonic_ns(),'window':int(wid),'error':repr(stale_error),'authority_granted':False})+'\n')
                continue
            with (case/'window-samples.jsonl').open('a') as sample:
                sample.write(json.dumps({'ns':time.monotonic_ns(),'window':int(wid),'legacy_title':str(legacy) if legacy is not None else None,'portable_title':title})+'\n')
            if isinstance(title,bytes):title=title.decode('utf-8','replace')
            if workbook.name in title:target=int(wid);break
        if time.monotonic()>deadline:raise RuntimeError('Calc window not ready')
        time.sleep(.05)
    if phase=='subsequent_reuse':
        subprocess.run(['wmctrl','-ir',hex(target),'-b','remove,maximized_vert,maximized_horz'],env=env,check=True)
        position=connection.screen().root.translate_coords(connection.create_resource_object('window',target),0,0)
        subprocess.run(['wmctrl','-ir',hex(target),'-e',f'0,{position.x+40},{position.y+30},-1,-1'],env=env,check=True)
    subprocess.run(['wmctrl','-ia',hex(target)],env=env,check=True)
    save('targets.json',{'app':target})
    from runtime.cli_v1.mcp_session import MCPSessionOwner
    from runtime.cli_v1.observe import observe_in_session
    from runtime.cli_v1.api import dispatch_in_session
    from runtime.cli_v1.review import present_result
    from PIL import Image
    capture_state=DeliveredCapture()
    owner=MCPSessionOwner({'app':target},f':{n}')
    save('owner.json',{'pid':os.getpid(),'display':f':{n}','window':target,'children':{name:p.pid for name,p in children},'started_ns':time.monotonic_ns()})
    (case/'commands').mkdir();(case/'replies').mkdir();(case/'public').mkdir()
    compact=sys.argv[4]=='compact';index=1;sequence=0;last_native=None
    print('READY public persistent API',flush=True)
    stop_at=time.monotonic()+900
    while True:
        if time.monotonic()>stop_at:raise RuntimeError('bounded owner lifetime exhausted')
        if any(p.poll() is not None for name,p in children):raise RuntimeError('owned child exited')
        path=case/'commands'/f'{index:03d}.json'
        if not path.exists():time.sleep(.02);continue
        request=json.loads(path.read_text());capture_region=region_for(request);started=time.monotonic_ns();op=request['op']
        if index>max_commands:raise RuntimeError('allocated command budget exhausted')
        projection=None
        if op=='observe':
            raw=observe_in_session(owner.get(),target='app',frame='screen_physical_px',region=capture_region,capture_directory=str(case/'public/images'))
            projection=present_result(raw,case/'public',compact=compact,report_refs=compact)
            last_native=capture_state.accept(raw,projection)
            if last_native is not None:sequence+=1
        elif op=='dispatch':
            if request['program']['source']!={'observation_seq':sequence,'binding_revision':owner.binding_revision}:raise ValueError('explicit current source and binding required')
            owner.dispatch_attempted=True
            raw=dispatch_in_session(owner.get(),request['program'],current_observation_seq=sequence,current_binding_revision=owner.binding_revision,capture_directory=str(case/'public/images'))
            raw['post_dispatch_inspection']=owner.inspect_after_dispatch(raw,'app',screen_region=capture_region,capture_directory=str(case/'public/images'),wait_ms=wait_ms)
            projection=present_result(raw,case/'public',compact=compact,report_refs=compact)
            last_native=capture_state.accept(raw,projection)
            if last_native is not None:sequence+=1
        elif op=='review_target':
            raw=owner.review_target(**request['arguments'],capture_directory=str(case/'public/images'))
            observed=raw.get('observation_report',{})
            if raw.get('capture_consistency')=='matched':
                projection=present_result(observed,case/'public',compact=compact,report_refs=compact)
            else:
                projection={'image':None,'image_status':'needs_review','authority':'none'}
            last_native=capture_state.accept(observed,projection)
            if last_native is not None:sequence+=1
        elif op=='close':raw=owner.close()
        else:raise ValueError('unsupported explicit public command')
        save(f'public/{index:03d}-raw.json',raw)
        save(f'replies/{index:03d}.json',{'request':request,'started_ns':started,'ended_ns':time.monotonic_ns(),'sequence':sequence,'binding_revision':owner.binding_revision,'raw_report':str(case/f'public/{index:03d}-raw.json'),'reply':projection if projection is not None else raw})
        print('REPLY '+str(index)+' '+op,flush=True);index+=1
        if op=='close':break
    owner.close();owner=None;host_exit=0

except Exception as error:
    save('exception.json',{'error':repr(error),'replay_allowed':False});raise
finally:
    if owner is not None:owner.close()
    if bridge is not None:bridge.close()
    if connection is not None:connection.close()
    codes=[]
    for name,p in reversed(children):
        if p.poll() is None:
            if name=='primary':p.terminate()
            else:os.killpg(p.pid,signal.SIGTERM)
        try:code=p.wait(timeout=5)
        except subprocess.TimeoutExpired:
            if name=='primary':p.kill()
            else:os.killpg(p.pid,signal.SIGKILL)
            code=p.wait(timeout=5)
        codes.append({'name':name,'pid':p.pid,'returncode':code})
    save('cleanup.json',{'host_exit':host_exit,'children':codes,'ended_ns':time.monotonic_ns()})
    # No application/file oracle is available to the controller before terminal.
    if 'workbook' in locals() and workbook.exists():
        book=load_workbook(workbook,read_only=True,data_only=True)
        cells={cell.coordinate:cell.value for row in book.active for cell in row if cell.value is not None}
        book.close();save('evaluation.json',{'expected':{'A1':expected[0],'A2':expected[1]},'actual_nonempty_cells':cells,'success':cells=={'A1':expected[0],'A2':expected[1]},'after_all_owned_processes_terminal':True,'workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    print('TERMINAL OWNER',flush=True)



