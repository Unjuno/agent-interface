#!/usr/bin/env python3
import argparse, csv, json, os, random, re, signal, subprocess, tempfile, time, xml.etree.ElementTree as ET
from pathlib import Path
from statistics import median

from PIL import ImageGrab
from Xlib import X, XK, display as xdisplay
from Xlib.ext import xtest
from openpyxl import Workbook, load_workbook

W,H=1280,800
PIXELS=W*H
MODE='sequential_full'
PRESS_DWELL_MS=0
CHAR_GAP_MS=0

class XSession:
    def __init__(self):
        self.tmp=Path(tempfile.mkdtemp(prefix='realapp-x-'))
        self.display_num=self._pick_display()
        self.name=f':{self.display_num}'
        self.auth=self.tmp/'Xauthority'; self.auth.touch()
        self.env=os.environ.copy(); self.env['DISPLAY']=self.name; self.env['XAUTHORITY']=str(self.auth)
        self.procs=[]
        self.xvfb=self._popen(['Xvfb',self.name,'-screen','0',f'{W}x{H}x24','-ac','-nolisten','tcp'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        sock=Path(f'/tmp/.X11-unix/X{self.display_num}')
        self._wait(lambda:sock.exists(),2.0,'Xvfb')
        self.openbox=self._popen(['openbox','--config-file','/etc/xdg/openbox/rc.xml'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        time.sleep(.20)
        os.environ['DISPLAY']=self.name; os.environ['XAUTHORITY']=str(self.auth)
        self.d=xdisplay.Display(self.name)
    def _pick_display(self):
        for n in range(120,190):
            if not Path(f'/tmp/.X11-unix/X{n}').exists() and not Path(f'/tmp/.X{n}-lock').exists(): return n
        raise RuntimeError('no display')
    def _popen(self,args,**kw):
        p=subprocess.Popen(args,env=self.env,start_new_session=True,**kw);self.procs.append(p);return p
    def spawn(self,args,**kw): return self._popen(args,**kw)
    def _wait(self,fn,timeout,label):
        end=time.perf_counter()+timeout
        while time.perf_counter()<end:
            if fn(): return
            time.sleep(.01)
        raise TimeoutError(label)
    def windows(self):
        p=subprocess.run(['wmctrl','-l'],env=self.env,text=True,capture_output=True)
        return p.stdout
    def wait_window(self,needle,timeout=8.0):
        t=time.perf_counter()
        self._wait(lambda: needle in self.windows(),timeout,f'window {needle}')
        return (time.perf_counter()-t)*1000
    def focus(self,needle=None,chromium=False):
        lines=self.windows().splitlines()
        win=None
        if chromium:
            for line in lines:
                if 'Chromium' in line: win=line.split()[0]; break
        elif needle:
            for line in lines:
                if needle in line: win=line.split()[0]; break
        if not win: raise RuntimeError(f'window not found {needle}: {lines}')
        subprocess.run(['wmctrl','-ia',win],env=self.env,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        time.sleep(.05)
    def close(self):
        try:self.d.close()
        except Exception:pass
        for p in reversed(self.procs):
            try: os.killpg(p.pid,signal.SIGTERM)
            except Exception: pass
        time.sleep(.05)
        for p in reversed(self.procs):
            try:
                if p.poll() is None: os.killpg(p.pid,signal.SIGKILL)
            except Exception: pass

class Driver:
    def __init__(self,sess,settle_ms=50):
        self.s=sess;self.d=sess.d;self.settle=settle_ms/1000
        self.input_events=0;self.logical_ops=0;self.obs=0;self.obs_pixels=0;self.capture_ms=[];self.pointer_path=0.0
    def _kc(self,name):
        ks=XK.string_to_keysym(name); kc=self.d.keysym_to_keycode(ks)
        if not kc: raise RuntimeError(f'no keycode: {name}')
        return kc
    def _raw(self,name,down):
        xtest.fake_input(self.d,X.KeyPress if down else X.KeyRelease,self._kc(name));self.d.sync();self.input_events+=1
    def key(self,name):
        self._raw(name,True);self._raw(name,False);self.logical_ops+=1;self.observe_after()
    def chord(self,mod,key):
        self._raw(mod,True);self._raw(key,True);self._raw(key,False);self._raw(mod,False);self.logical_ops+=1;self.observe_after()
    def _char(self,ch):
        if ch.isalpha() or ch.isdigit():
            self._raw(ch,True);self._raw(ch,False);return
        mapping={':':('semicolon',True),'/':('slash',False),'-':('minus',False),'.':('period',False),'_':('minus',True),' ':('space',False)}
        name,shift=mapping[ch]
        if shift:self._raw('Shift_L',True)
        self._raw(name,True);self._raw(name,False)
        if shift:self._raw('Shift_L',False)
    def text(self,text):
        for i,ch in enumerate(text):
            self._char(ch)
            if CHAR_GAP_MS and i+1<len(text): time.sleep(CHAR_GAP_MS/1000)
        self.logical_ops+=1;self.observe_after()
    def drag(self,x0,y0,x1,y1):
        xtest.fake_input(self.d,X.MotionNotify,x=x0,y=y0);self.d.sync();self.input_events+=1
        xtest.fake_input(self.d,X.ButtonPress,1);self.d.sync();self.input_events+=1
        if PRESS_DWELL_MS: time.sleep(PRESS_DWELL_MS/1000)
        steps=max(4,min(12,int((((x1-x0)**2+(y1-y0)**2)**.5)//4)))
        for i in range(1,steps+1):
            x=round(x0+(x1-x0)*i/steps); y=round(y0+(y1-y0)*i/steps)
            xtest.fake_input(self.d,X.MotionNotify,x=x,y=y);self.d.sync();self.input_events+=1;time.sleep(.004)
        self.pointer_path+=((x1-x0)**2+(y1-y0)**2)**.5
        xtest.fake_input(self.d,X.ButtonRelease,1);self.d.sync();self.input_events+=1
        self.logical_ops+=1;self.observe_after()
    def capture(self):
        t=time.perf_counter(); im=ImageGrab.grab(); ms=(time.perf_counter()-t)*1000
        self.capture_ms.append(ms);self.obs+=1;self.obs_pixels+=im.width*im.height
        return im
    def initial_observe(self): return self.capture()
    def observe_after(self):
        if MODE!='sequential_full': return None
        if self.settle:time.sleep(self.settle)
        return self.capture()

def wait_until(fn,timeout=2.0,step=.01):
    end=time.perf_counter()+timeout
    while time.perf_counter()<end:
        try:
            if fn():return True
        except Exception:pass
        time.sleep(step)
    return False

def largest_red_bbox(im):
    # Restrict to document workspace, avoiding red icons in toolbars.
    pix=im.convert('RGB'); xs=[];ys=[]
    for y in range(180,700,2):
        for x in range(180,820,2):
            r,g,b=pix.getpixel((x,y))
            if r>180 and g<110 and b<110:
                xs.append(x);ys.append(y)
    if len(xs)<100: raise RuntimeError(f'red target not found n={len(xs)}')
    # Our benchmark object is the dominant dense red region. Robust quantiles reject stray icons.
    xs.sort();ys.sort();q=lambda a,p:a[int((len(a)-1)*p)]
    return (q(xs,.05),q(ys,.05),q(xs,.95),q(ys,.95))

def task_xterm(seed,settle):
    s=XSession();tmp=s.tmp; token=f't{seed:04d}'; out=tmp/'out.txt'
    try:
        t0=time.perf_counter()
        s.spawn(['xterm','-T','AI-XTERM','-geometry','80x24','-e','sh','-c',f'IFS= read -r line; printf "%s" "$line" > {out}; sleep .3'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        ready=s.wait_window('AI-XTERM');s.focus('AI-XTERM'); launch=(time.perf_counter()-t0)*1000
        d=Driver(s,settle);d.initial_observe();t=time.perf_counter();d.text(token);d.key('Return')
        ok=wait_until(lambda:out.exists() and out.read_text()==token,1.0);task=(time.perf_counter()-t)*1000
        return record('xterm',seed,ok,launch,task,d,{'token_len':len(token)})
    finally:s.close()

def task_chromium(seed,settle):
    s=XSession();tmp=s.tmp
    try:
        t0=time.perf_counter();s.spawn(['chromium','--no-sandbox','--disable-gpu','--disable-dev-shm-usage','--no-first-run','--no-default-browser-check',f'--user-data-dir={tmp}/profile','about:blank'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        ready=s.wait_window('Chromium',10);s.focus(chromium=True);launch=(time.perf_counter()-t0)*1000
        d=Driver(s,settle);d.initial_observe();t=time.perf_counter();d.chord('Control_L','l');d.text('chrome://version');d.key('Return')
        ok=wait_until(lambda:'chrome://version/' in s.windows(),2.0);task=(time.perf_counter()-t)*1000
        return record('chromium',seed,ok,launch,task,d,{})
    finally:s.close()

def task_calc(seed,settle):
    s=XSession();tmp=s.tmp; path=tmp/'sheet.xlsx';rng=random.Random(seed);a=rng.randint(100,899);b=rng.randint(100,899)
    wb=Workbook();wb.save(path);before=path.stat().st_mtime_ns
    try:
        t0=time.perf_counter();s.spawn(['libreoffice','--norestore','--nodefault','--nolockcheck',f'-env:UserInstallation=file://{tmp}/loprofile','--calc',str(path)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        ready=s.wait_window('sheet.xlsx',12);s.focus('sheet.xlsx');time.sleep(.2);launch=(time.perf_counter()-t0)*1000
        d=Driver(s,settle);d.initial_observe();t=time.perf_counter();d.text(str(a));d.key('Return');d.text(str(b));
        if MODE=='sparse_reactive':
            before_windows=s.windows(); d.chord('Control_L','s')
            modal_ok=wait_until(lambda:'Confirm File Format' in s.windows(),1.0,.002)
            if modal_ok: d.capture()
            d.key('Return')
        else:
            d.chord('Control_L','s');d.key('Return')
        def verify():
            if path.stat().st_mtime_ns==before:return False
            try:
                w=load_workbook(path,data_only=False,read_only=True);sh=w.active;v=(sh['A1'].value,sh['A2'].value);w.close();return v==(a,b)
            except Exception:return False
        ok=wait_until(verify,2.5);task=(time.perf_counter()-t)*1000
        return record('calc',seed,ok,launch,task,d,{'a':a,'b':b})
    finally:s.close()

def task_inkscape(seed,settle):
    s=XSession();tmp=s.tmp;path=tmp/'shape.svg';path.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200" viewBox="0 0 200 200"><rect id="r" x="50" y="50" width="40" height="30" fill="red"/></svg>')
    rng=random.Random(seed);dx=rng.choice([24,30,36])
    try:
        t0=time.perf_counter();s.spawn(['inkscape',str(path)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        ready=s.wait_window('shape.svg',10);s.focus('shape.svg');time.sleep(.9);launch=(time.perf_counter()-t0)*1000
        d=Driver(s,settle);im=d.initial_observe();bbox=largest_red_bbox(im);x0=(bbox[0]+bbox[2])//2;y0=(bbox[1]+bbox[3])//2
        t=time.perf_counter();d.chord('Control_L','a')
        if MODE=='sparse_reactive': time.sleep(.10)
        d.drag(x0,y0,x0+dx,y0);d.chord('Control_L','s')
        def xpos():
            try:
                root=ET.parse(path).getroot();r=root.find('{http://www.w3.org/2000/svg}rect');return float(r.attrib.get('x','50'))
            except Exception:return 50.0
        ok=wait_until(lambda:xpos()>50.5,2.0);task=(time.perf_counter()-t)*1000
        return record('inkscape',seed,ok,launch,task,d,{'dx_screen':dx,'x_svg':xpos(),'bbox':bbox})
    finally:s.close()

def record(app,seed,ok,launch,task,d,extra):
    return {'app':app,'seed':seed,'success':bool(ok),'launch_ms':launch,'task_ms':task,'logical_ops':d.logical_ops,'input_events':d.input_events,'observations':d.obs,'observed_mpix':d.obs_pixels/1e6,'capture_p50_ms':median(d.capture_ms) if d.capture_ms else 0,'capture_max_ms':max(d.capture_ms) if d.capture_ms else 0,'pointer_path_px':d.pointer_path,'extra':extra}

TASKS={'xterm':task_xterm,'chromium':task_chromium,'calc':task_calc,'inkscape':task_inkscape}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--app',choices=TASKS);ap.add_argument('--episodes',type=int,default=1);ap.add_argument('--seed0',type=int,default=1);ap.add_argument('--settle-ms',type=int,default=50);ap.add_argument('--mode',choices=['sequential_full','sparse_reactive'],default='sequential_full');ap.add_argument('--press-dwell-ms',type=float,default=0);ap.add_argument('--char-gap-ms',type=float,default=0);ap.add_argument('--out');a=ap.parse_args();
    global MODE,PRESS_DWELL_MS,CHAR_GAP_MS;MODE=a.mode;PRESS_DWELL_MS=a.press_dwell_ms;CHAR_GAP_MS=a.char_gap_ms
    rows=[]
    for i in range(a.episodes):
        seed=a.seed0+i
        try:r=TASKS[a.app](seed,a.settle_ms)
        except Exception as e:r={'app':a.app,'seed':seed,'success':False,'error':repr(e)}
        rows.append(r);print(json.dumps(r,sort_keys=True),flush=True)
    if a.out:Path(a.out).write_text('\n'.join(json.dumps(r,sort_keys=True) for r in rows)+'\n')

if __name__=='__main__':main()
