#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json, os, subprocess, time
from pathlib import Path
from openpyxl import Workbook
from Xlib import display
HERE=Path(__file__).resolve().parent

def wait_display(env,name):
    for _ in range(100):
        if subprocess.run(['xdpyinfo','-display',name],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0:return
        time.sleep(.04)
    raise RuntimeError('Xvfb not ready')

def find_calc(display_name,timeout=8.0):
    deadline=time.monotonic()+timeout
    while time.monotonic()<deadline:
        d=display.Display(display_name)
        def walk(w):
            try: kids=w.query_tree().children
            except Exception:return None
            for c in kids:
                try: cls=c.get_wm_class()
                except Exception: cls=None
                if cls and any('libreoffice-calc' in str(x).lower() for x in cls): return c.id
                q=walk(c)
                if q is not None:return q
            return None
        found=walk(d.screen().root); d.close()
        if found is not None:return found
        time.sleep(.05)
    raise RuntimeError('Calc window not found')

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--out',type=Path,required=True); ap.add_argument('--display',required=True); ap.add_argument('--pacing-ms',type=float,required=True); ap.add_argument('--controller',type=Path,required=True); a=ap.parse_args(); a.out.mkdir(parents=True,exist_ok=False)
    xlsx=a.out/'task.xlsx'; wb=Workbook(); ws=wb.active; ws['A1']='seed'; ws['A2']='old'; wb.save(xlsx); input_sha=hashlib.sha256(xlsx.read_bytes()).hexdigest()
    auth=a.out/'Xauthority'; auth.write_bytes(b''); env=os.environ.copy(); env['DISPLAY']=a.display; env['XAUTHORITY']=str(auth); os.environ['DISPLAY']=a.display; os.environ['XAUTHORITY']=str(auth)
    xvfb=subprocess.Popen(['Xvfb',a.display,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(a.out/'xvfb.stdout').open('w'),stderr=(a.out/'xvfb.stderr').open('w')); ob=lo=None
    try:
        wait_display(env,a.display); ob=subprocess.Popen(['openbox'],env=env,stdout=(a.out/'openbox.stdout').open('w'),stderr=(a.out/'openbox.stderr').open('w'))
        lo=subprocess.Popen(['libreoffice',f'-env:UserInstallation=file://{a.out}/profile','--nologo','--nodefault','--nofirststartwizard','--norestore',str(xlsx)],env=env,stdout=(a.out/'libreoffice.stdout').open('w'),stderr=(a.out/'libreoffice.stderr').open('w'))
        wid=find_calc(a.display); time.sleep(.4)
        exec_json=a.out/'execution.json'
        run=subprocess.run([str(a.controller),'--display',a.display,'--window-id',str(wid),'--pacing-ms',str(a.pacing_ms),'--out',str(exec_json)],env=env,text=True,capture_output=True)
        (a.out/'controller.stdout').write_text(run.stdout); (a.out/'controller.stderr').write_text(run.stderr); (a.out/'controller.exitcode').write_text(str(run.returncode)+'\n')
        time.sleep(.8)
        score=subprocess.run(['python3',str(HERE/'score_workbook.py'),'--xlsx',str(xlsx),'--out',str(a.out/'score.json')],text=True,capture_output=True)
        (a.out/'scorer.stdout').write_text(score.stdout); (a.out/'scorer.stderr').write_text(score.stderr); (a.out/'scorer.exitcode').write_text(str(score.returncode)+'\n')
        ex=json.loads(exec_json.read_text()); sc=json.loads((a.out/'score.json').read_text())
        report={'schema':'agent-interface/native-x11-text-pacing-arm-v1','pacing_ms':a.pacing_ms,'window_id':wid,'input_xlsx_sha256':input_sha,'controller_exitcode':run.returncode,'scorer_exitcode':score.returncode,'transport_passed':ex.get('passed_transport') is True,'release_verified':bool(ex.get('task',{}).get('final_release_verified') and ex.get('confirm',{}).get('final_release_verified')),'stale_zero_injected_events':(ex.get('stale',{}).get('accepted') is False and ex.get('stale',{}).get('error')=='STALE_OBSERVATION' and ex.get('stale',{}).get('injected_events')==0),'exact_count':sc['exact_count'],'semantic_passed':sc['passed'],'edit_elapsed_ns':ex['edit_elapsed_ns'],'median_char_call_ns':ex['median_char_call_ns'],'median_char_start_interval_ns':ex['median_char_start_interval_ns'],'output_xlsx_sha256':sc['xlsx_sha256']}
        report['eligible']=all([report['transport_passed'],report['release_verified'],report['stale_zero_injected_events'],report['semantic_passed']])
        (a.out/'report.json').write_text(json.dumps(report,indent=2)+'\n'); print(json.dumps(report,indent=2)); return 0 if report['eligible'] else 1
    finally:
        if lo and lo.poll() is None: lo.terminate()
        if lo:
            try: lo.wait(timeout=3)
            except subprocess.TimeoutExpired: lo.kill()
        if ob and ob.poll() is None: ob.terminate()
        xvfb.terminate()
if __name__=='__main__': raise SystemExit(main())
