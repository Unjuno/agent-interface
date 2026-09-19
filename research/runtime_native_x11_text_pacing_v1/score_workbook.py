#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook
CORPUS=['office','coffee','bookkeeper','committee','parallel','address','success','letterpress','mississippi','assessment','committee','bookkeeping','ffffffffff','aaaaaaaaaa','tttttttttt','ssssssssss']
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--xlsx',type=Path,required=True); ap.add_argument('--out',type=Path,required=True); a=ap.parse_args()
    wb=load_workbook(a.xlsx,data_only=False); ws=wb.active; got=[ws.cell(i+1,1).value for i in range(len(CORPUS))]
    mism=[{'row':i+1,'expected':e,'actual':g} for i,(e,g) in enumerate(zip(CORPUS,got)) if e!=g]
    result={'schema':'agent-interface/native-x11-text-pacing-score-v1','expected':CORPUS,'actual':got,'mismatches':mism,'exact_count':len(CORPUS)-len(mism),'passed':not mism,'xlsx_sha256':hashlib.sha256(a.xlsx.read_bytes()).hexdigest(),'xlsx_bytes':a.xlsx.stat().st_size}
    a.out.write_text(json.dumps(result,indent=2)+'\n'); print(json.dumps(result,indent=2)); return 0 if result['passed'] else 1
if __name__=='__main__': raise SystemExit(main())
