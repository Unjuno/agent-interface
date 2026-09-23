#!/usr/bin/env python3
from pathlib import Path
import os, subprocess, time, json, hashlib, shutil, stat
from openpyxl import Workbook
from Xlib import display, X
from office_backend import OfficeX11Backend

ROOT=Path('/mnt/data/libreoffice_conflict_c254/scored-pair-03')
if ROOT.exists(): shutil.rmtree(ROOT)
ROOT.mkdir(parents=True)

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def statrow(p):
    s=Path(p).stat(); return {'dev':s.st_dev,'ino':s.st_ino,'size':s.st_size,'mtime_ns':s.st_mtime_ns,'ctime_ns':s.st_ctime_ns,'sha256':sha(p)}
def make_xlsx(p,a1,a2,a3=None,b1=None):
    wb=Workbook(); ws=wb.active; ws['A1']=a1; ws['A2']=a2; ws['A3']=a3; ws['B1']=b1; wb.save(p)
def wait_display(env,name):
    for _ in range(100):
        if subprocess.run(['xdpyinfo','-display',name],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0:return
        time.sleep(.05)
    raise RuntimeError('Xvfb not ready')
def windows(display_name):
    d=display.Display(display_name); root=d.screen().root; out=[]
    def walk(w,depth=0):
        try:
            name=w.get_wm_name(); cls=w.get_wm_class(); geo=w.get_geometry(); attrs=w.get_attributes()
        except Exception:
            name=cls=geo=attrs=None
        if name or cls:
            out.append({'id':w.id,'name':str(name) if name else None,'class':list(cls) if cls else None,'depth':depth,
                        'w':getattr(geo,'width',None) if geo else None,'h':getattr(geo,'height',None) if geo else None,
                        'map_state':getattr(attrs,'map_state',None) if attrs else None})
        try: children=w.query_tree().children
        except Exception: children=[]
        for c in children: walk(c,depth+1)
    walk(root); d.close(); return out
def find_calc(display_name):
    deadline=time.monotonic()+10
    while time.monotonic()<deadline:
        for r in windows(display_name):
            if r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']): return r['id']
        time.sleep(.05)
    raise RuntimeError('Calc window not found')
def modal_names(rows,calc_id):
    names=[]
    for r in rows:
        if r['id']==calc_id: continue
        if r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']) and r['name']:
            names.append(r['name'])
    return sorted(set(names))

def run_arm(name,display_name,mode):
    out=ROOT/name; out.mkdir()
    xlsx=out/'task.xlsx'; make_xlsx(xlsx,'seed','old')
    initial=statrow(xlsx)
    auth=out/'Xauthority'; auth.write_bytes(b'')
    env=os.environ.copy(); env['DISPLAY']=display_name; env['XAUTHORITY']=str(auth)
    os.environ['DISPLAY']=display_name; os.environ['XAUTHORITY']=str(auth)
    xvfb=subprocess.Popen(['Xvfb',display_name,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(out/'xvfb.stdout').open('w'),stderr=(out/'xvfb.stderr').open('w'))
    ob=lo=None; b=None
    result={'arm':name,'mode':mode,'display':display_name,'initial_file':initial}
    try:
        wait_display(env,display_name)
        ob=subprocess.Popen(['openbox'],env=env,stdout=(out/'openbox.stdout').open('w'),stderr=(out/'openbox.stderr').open('w'))
        profile=f'file://{out}/profile'
        lo=subprocess.Popen(['libreoffice',f'-env:UserInstallation={profile}','--nologo','--nodefault','--nofirststartwizard','--norestore',str(xlsx)],env=env,stdout=(out/'libreoffice.stdout').open('w'),stderr=(out/'libreoffice.stderr').open('w'))
        calc=find_calc(display_name); result['calc_window_id']=calc; time.sleep(.5)
        b=OfficeX11Backend(display_name,{'calc':calc})
        b.focus('calc'); b.pointer_move('calc','window_client',80,180); b.pointer_button('left',True); b.pointer_button('left',False)
        b.key_chord(['CTRL','Home']); b.text('office'); b.key_chord(['ENTER']); b.text('preview'); b.key_chord(['ENTER']); time.sleep(.15)
        d=display.Display(display_name); focus=d.get_input_focus().focus; focus_id=getattr(focus,'id',None); d.close()
        result['precheck']={'focus_id':focus_id,'calc_window_id':calc,'windows':windows(display_name),'file':statrow(xlsx),'monotonic_ns':time.monotonic_ns()}
        if mode=='inplace':
            repl=out/'replacement.xlsx'; make_xlsx(repl,'external','replacement','external-marker','writer')
            result['replacement_source']=statrow(repl)
            before_inplace=statrow(xlsx); payload=repl.read_bytes()
            with xlsx.open('r+b') as f:
                f.seek(0); f.write(payload); f.truncate(); f.flush(); os.fsync(f.fileno())
            after_inplace=statrow(xlsx)
            if after_inplace['ino'] != before_inplace['ino']:
                raise RuntimeError('in-place mutation changed inode')
            result['after_external_inplace']={'file':after_inplace,'pre_inode':before_inplace['ino'],'monotonic_ns':time.monotonic_ns()}
        result['before_save']={'file':statrow(xlsx),'windows':windows(display_name),'monotonic_ns':time.monotonic_ns()}
        b.key_chord(['CTRL','S']); time.sleep(.7)
        first=windows(display_name); first_modals=modal_names(first,calc)
        result['after_ctrl_s']={'windows':first,'modal_names':first_modals,'file':statrow(xlsx),'monotonic_ns':time.monotonic_ns()}
        if first_modals == ['Confirm File Format']:
            b.key_chord(['ENTER']); result['format_confirm_enter_sent']=True; time.sleep(1.3)
        else:
            result['format_confirm_enter_sent']=False
        second=windows(display_name); second_modals=modal_names(second,calc)
        result['after_optional_format_confirm']={'windows':second,'modal_names':second_modals,'file':statrow(xlsx),'monotonic_ns':time.monotonic_ns()}
        rel=b.release_all().__dict__; result['release']=rel; result['backend_emissions']=b.emissions
        score_path=out/'score.json'
        sp=subprocess.run(['python','/mnt/data/libreoffice_conflict_c254/score_any.py','--xlsx',str(xlsx),'--out',str(score_path)],capture_output=True,text=True)
        (out/'scorer.stdout').write_text(sp.stdout); (out/'scorer.stderr').write_text(sp.stderr); result['scorer_exitcode']=sp.returncode
        result['score']=json.loads(score_path.read_text())
        cells=result['score'].get('cells')
        if mode=='stable':
            result['classification']='stable_saved' if cells and cells.get('A1')=='office' and cells.get('A2')=='preview' else 'stable_other'
        else:
            if cells and cells.get('A1')=='external' and cells.get('A2')=='replacement' and cells.get('A3')=='external-marker':
                result['classification']='inplace_external_preserved_with_prompt' if (first_modals!=['Confirm File Format'] or second_modals) else 'inplace_external_preserved_no_conflict_prompt'
            elif cells and cells.get('A1')=='office' and cells.get('A2')=='preview':
                result['classification']='stale_overwrite'
            else:
                result['classification']='uncertain_other'
        (out/'result.json').write_text(json.dumps(result,indent=2,default=str)+'\n')
        return result
    finally:
        if b:
            try: b.release_all(); b.close()
            except Exception: pass
        if lo and lo.poll() is None: lo.terminate()
        if lo:
            try: lo.wait(timeout=3)
            except Exception: lo.kill()
        if ob and ob.poll() is None: ob.terminate()
        xvfb.terminate()

rows=[]
for spec in [('stable',':223','stable'),('inplace',':224','inplace')]:
    rows.append(run_arm(*spec))
summary={'schema':'agent-interface/libreoffice-save-conflict-inplace-pair-v1','allocation':'c254-pair-03','order':['stable','inplace'],'rows':rows,
         'source_blobs':{'backend_x11.py':'b4f8e043ce4f8929d446e038418ea0fd3655bab0','office_backend.py':'3aeca10f62fb1366bcdd5fad561509cfcc0d91ab'}}
(ROOT/'SUMMARY.json').write_text(json.dumps(summary,indent=2,default=str)+'\n')
print(json.dumps({'allocation':summary['allocation'],'outcomes':[{'arm':r['arm'],'classification':r['classification'],'modals1':r['after_ctrl_s']['modal_names'],'modals2':r['after_optional_format_confirm']['modal_names'],'score':r['score']} for r in rows]},indent=2))
