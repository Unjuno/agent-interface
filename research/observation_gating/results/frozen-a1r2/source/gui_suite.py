#!/usr/bin/env python3
"""O0/O1 live X11 experiment; model-boundary images are a local sink, not an LLM.

The controller only sees Receiver output and public X11 state. File/HTTP output
oracles run after it has returned. Images are archived and audited after timing.
"""

import argparse
import hashlib
import http.server
import json
import os
import platform
import random
import shutil
import signal
import subprocess
import sys
import threading
import time
import traceback
import urllib.parse
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np
from PIL import Image, ImageGrab
from openpyxl import Workbook, load_workbook

from exact_gate import ExactGate, Frame, Receiver

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "real_apps_v1"))
import real_app_suite_v1 as base

base.MODE = "sparse_reactive"
base.CHAR_GAP_MS = 2.0
base.PRESS_DWELL_MS = 5.0
APPS = ("xterm", "chromium", "calc", "inkscape")
PARAMETERS = {
    "display": "Xvfb 1280x800x24 + Openbox + XTEST",
    "reset": "fresh X server, WM, application process, HOME, XDG and profile per episode",
    "char_gap_ms": 2.0, "press_dwell_ms": 5.0, "motion_gap_ms": 4.0,
    "poll_ms": 16.0, "public_timeout_s": 4.0, "oracle_timeout_s": 3.0,
    "gate": "exact dimensions + mode + immutable pixel bytes; no hash/threshold",
    "first_feedback": "fresh image or timestamped unchanged-image reference with public X11 context",
    "first_feedback_is_semantic_completion": False,
    "calc_setup_readiness": "painted worksheet: >20% near-white pixels and >128 colors; setup only",
    "inkscape_selection_barrier": "visible dark selector handles on all four sides of acquired red object",
    "inkscape_setup_readiness": "visible red document target before task timing",
    "model_calls": 0,
}


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True) + "\n")


def code_hashes():
    paths = [HERE / name for name in ("exact_gate.py", "gui_suite.py", "test_exact_gate.py",
                                    "analyze.py", "PROTOCOL.md")]
    paths.append(HERE.parent / "real_apps_v1" / "real_app_suite_v1.py")
    return {str(p.relative_to(HERE.parent)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in paths}


def environment(chromium):
    versions = {}
    diagnostics = {}
    for name, command in {
        "xterm": ["xterm", "-version"], "chromium": [chromium, "--version"],
        "calc": ["libreoffice", "--version"], "inkscape": ["inkscape", "--version"],
        "openbox": ["openbox", "--version"],
        "packages": ["dpkg-query", "-W", "xvfb", "python3-xlib", "python3-pil",
                     "python3-numpy", "python3-openpyxl"],
    }.items():
        result = subprocess.run(command, capture_output=True, text=True, timeout=20)
        versions[name] = result.stdout.strip() or result.stderr.strip()
        diagnostics[name] = result.stderr.strip()
    return {"platform": platform.platform(), "python": sys.version,
            "os_release": Path("/etc/os-release").read_text(), "versions": versions,
            "version_diagnostics": diagnostics,
            "chromium_path": chromium, "parameters": PARAMETERS,
            "code_sha256": code_hashes(), "utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}


class Session(base.XSession):
    def __init__(self):
        try:
            super().__init__()
        except Exception:
            self.close()
            raise
        # Keep all application state private; explicitly capture this X server.
        for key in ("WAYLAND_DISPLAY", "DBUS_SESSION_BUS_ADDRESS"):
            self.env.pop(key, None)
        for key, folder in (("HOME", "home"), ("XDG_CONFIG_HOME", "config"),
                            ("XDG_CACHE_HOME", "cache"), ("XDG_DATA_HOME", "data"),
                            ("XDG_RUNTIME_DIR", "run")):
            path = self.tmp / folder
            path.mkdir(mode=0o700)
            self.env[key] = str(path)
        self.env.update(GDK_BACKEND="x11", QT_QPA_PLATFORM="xcb", LANG="C.UTF-8",
                        LC_ALL="C.UTF-8", SAL_USE_VCLPLUGIN="gen")

    def _wait(self, fn, timeout, label):
        if label == "Xvfb":
            def connected():
                try:
                    connection = base.xdisplay.Display(self.name)
                    connection.close()
                    return True
                except Exception:
                    return False
            # WSLg can prevent a pathname socket while the abstract UNIX socket
            # is healthy. Readiness means a successful X11 handshake.
            fn = connected
        return super()._wait(fn, timeout, label)

    def context(self):
        atom = self.d.intern_atom("_NET_ACTIVE_WINDOW")
        prop = self.d.screen().root.get_full_property(atom, 0)
        focus = int(prop.value[0]) if prop is not None and len(prop.value) else 0
        return (("windows", self.windows()), ("active_window", focus))

    def close(self):
        super().close()
        for process in reversed(getattr(self, "procs", [])):
            try:
                process.wait(timeout=2)
            except subprocess.TimeoutExpired:
                pass


def image_for(frame):
    return Image.frombytes(frame.mode, (frame.width, frame.height), frame.pixels)


def red_bbox(frame):
    """Visible fixture-target acquisition, not a gate or saved-file oracle.

    Find the largest wide red run component without fixed app coordinates.
    Small palette cells and toolbar icons are excluded by component size.
    """
    arr = np.asarray(image_for(frame).convert("RGB"))
    mask = (arr[:, :, 0] > 180) & (arr[:, :, 1] < 100) & (arr[:, :, 2] < 100)
    candidates = []
    active = None
    for y, row in enumerate(mask):
        edges = np.flatnonzero(np.diff(np.r_[False, row, False].astype(np.int8)))
        spans = [(int(a), int(b)) for a, b in zip(edges[::2], edges[1::2]) if b - a >= 24]
        if spans:
            a, b = max(spans, key=lambda span: span[1] - span[0])
            if active is not None and a < active[2] and b > active[0]:
                active = (min(a, active[0]), active[1], max(b, active[2]), y + 1)
            else:
                if active is not None:
                    candidates.append(active)
                active = (a, y, b, y + 1)
        elif active is not None:
            candidates.append(active)
            active = None
    if active is not None:
        candidates.append(active)
    candidates = [b for b in candidates if b[3] - b[1] >= 16]
    return max(candidates, key=lambda b: (b[2] - b[0]) * (b[3] - b[1]), default=None)


def selection_visible(frame):
    """Public precondition before a dependent drag, shared by O0 and O1.

    Ctrl+A delivery is not proof that Inkscape consumed the selection command.
    Inspect the selector handles around the acquired object, with no fixed app
    coordinates and no saved-document access. This never suppresses an image.
    """
    bbox = red_bbox(frame)
    if bbox is None:
        return False
    x0, y0, x1, y1 = bbox
    arr = np.asarray(image_for(frame).convert("RGB"))
    dark = np.all(arr < 80, axis=2)
    left, right = max(0, x0 - 20), min(frame.width, x1 + 20)
    top, bottom = max(0, y0 - 20), min(frame.height, y1 + 20)
    bands = [dark[y0:y1, left:x0], dark[y0:y1, x1:right],
             dark[top:y0, x0:x1], dark[y1:bottom, x0:x1]]
    return all(np.count_nonzero(band) >= 5 for band in bands)


class Observer:
    def __init__(self, session, strategy, stream):
        self.session = session
        self.gate = ExactGate(strategy, stream)
        self.receiver = Receiver(stream)
        self.samples = []
        self.actions = []
        self.last = None
        self.last_context = ()

    def sample(self, action_id, reason):
        t0 = time.perf_counter_ns()
        image = ImageGrab.grab(xdisplay=self.session.name)
        tc = time.perf_counter_ns()
        # Mode and dimensions are part of exact identity; conversion is charged.
        image = image.convert("RGB")
        frame = Frame(image.width, image.height, image.mode, image.tobytes())
        ts = time.perf_counter_ns()
        context = self.session.context()
        tp = time.perf_counter_ns()
        update = self.gate.push(frame, observed_ns=tc, action_id=action_id, context=context)
        tg = time.perf_counter_ns()
        delivered = self.receiver.accept(update)
        tr = time.perf_counter_ns()
        self.samples.append({"source": frame, "delivered": delivered, "update": update,
                             "reason": reason, "capture_ns": tc - t0,
                             "serialize_ns": ts - tc, "context_ns": tp - ts,
                             "gate_ns": tg - tp, "receive_ns": tr - tg,
                             "ready_ns": tr, "sample_ns": tr - t0})
        self.last, self.last_context = delivered, update.context
        return delivered

    def action(self, label, operation, condition=None):
        issued = time.perf_counter_ns()
        operation()
        injected = time.perf_counter_ns()
        self.sample(label, "after_action")
        first = time.perf_counter_ns()
        deadline = injected + int(PARAMETERS["public_timeout_s"] * 1e9)
        public_known = None
        timed_out = False
        if condition is not None:
            while not condition(self.last, dict(self.last_context)):
                if time.perf_counter_ns() >= deadline:
                    timed_out = True
                    break
                # A sampling cadence, never a semantic-completion barrier.
                next_sample = self.samples[-1]["ready_ns"] + int(PARAMETERS["poll_ms"] * 1e6)
                time.sleep(max(0.0, (next_sample - time.perf_counter_ns()) / 1e9))
                self.sample(label, "waiting_public_effect")
            if not timed_out:
                public_known = time.perf_counter_ns()
        self.actions.append({"action_id": label, "issued_ns": issued,
                             "input_ack_ns": injected, "first_feedback_ns": first,
                             "public_effect_known_ns": public_known, "timed_out": timed_out,
                             "input_ack_ms": (injected - issued) / 1e6,
                             "first_feedback_ms": (first - issued) / 1e6,
                             "public_effect_ms": None if public_known is None else (public_known - issued) / 1e6})
        if timed_out:
            raise TimeoutError("Public effect: " + label)


def fixture_server(path):
    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def do_GET(self):
            body = (b'<!doctype html><title>AI FORM READY</title><h1>Observation fixture</h1>'
                    b'<form method="post" action="/submit"><label>Value '
                    b'<input name="value" autofocus></label><button>Save</button></form>')
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_POST(self):
            data = self.rfile.read(int(self.headers["Content-Length"]))
            # The app logs its output. It does not know the expected answer.
            path.write_bytes(data)
            body = b'<!doctype html><title>AI FORM SAVED</title><h1>Submission received</h1>'
            self.send_response(200)
            self.send_header("Content-Type", "text/html")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server


def prepare(session, app, seed, chromium):
    rng = random.Random(seed)
    goal = {"token": f"t{seed:06d}", "a": rng.randint(100, 899),
            "b": rng.randint(100, 899), "dx": rng.choice([24, 30, 36])}
    server = None
    if app == "xterm":
        output = session.tmp / "submitted.txt"
        script = session.tmp / "terminal_fixture.py"
        script.write_text("import pathlib,sys,signal\nline=input()\n"
                          "pathlib.Path(sys.argv[1]).write_text(line)\n"
                          "print('\\033]0;AI XTERM SAVED\\007',end='',flush=True)\n"
                          "signal.pause()\n")
        args = ["xterm", "-T", "AI XTERM READY", "-geometry", "80x24", "-fa", "DejaVu Sans Mono",
                "-e", sys.executable, str(script), str(output)]
        window = "AI XTERM READY"
    elif app == "chromium":
        output = session.tmp / "submitted.txt"
        server = fixture_server(output)
        goal["url"] = f"http://127.0.0.1:{server.server_port}/"
        args = [chromium, "--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage",
                "--no-first-run", "--no-default-browser-check", "--disable-background-networking",
                "--disable-component-update", "--disable-sync", "--password-store=basic",
                f"--user-data-dir={session.tmp}/browser-profile", "about:blank"]
        window = "about:blank"
    elif app == "calc":
        output = session.tmp / "sheet.xlsx"
        wb = Workbook()
        wb.save(output)
        profile = session.tmp / "lo-profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user" / "registrymodifications.xcu").write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="ShowTipOfTheDay" '
            'oor:op="fuse"><value>false</value></prop></item></oor:items>')
        args = ["libreoffice", "--norestore", "--nodefault", "--nolockcheck",
                f"-env:UserInstallation={profile.as_uri()}", "--calc", str(output)]
        window = "sheet.xlsx"
    else:
        output = session.tmp / "shape.svg"
        output.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200" '
                          'viewBox="0 0 200 200"><rect id="r" x="50" y="50" width="40" '
                          'height="30" fill="red"/></svg>')
        args = ["inkscape", str(output)]
        window = "shape.svg"
    with (session.tmp / "application.log").open("w") as log:
        session.spawn(args, stdout=log, stderr=log)
    session.wait_window(window, 20.0)
    session.focus(window)
    setup_frames = 0
    if app in ("calc", "inkscape"):
        # Window mapped != document painted. A blank grey client accepts input
        # but must not provide the apparent redundancy for this experiment.
        deadline = time.perf_counter() + 15
        while True:
            im = ImageGrab.grab(xdisplay=session.name).convert("RGB")
            setup_frames += 1
            arr = np.asarray(im)
            ready = (np.mean(np.all(arr > 245, axis=2)) > .2 and im.getcolors(128) is None
                     if app == "calc" else red_bbox(Frame(im.width, im.height, "RGB", im.tobytes())) is not None)
            if ready:
                break
            if time.perf_counter() >= deadline:
                raise TimeoutError("Painted document readiness: " + app)
            time.sleep(.016)
    goal["setup_readiness_captures"] = setup_frames
    return goal, output, server


def controller(session, app, goal, observer):
    """No output-path argument, no file oracle, and no raw captured-frame access."""
    driver = base.Driver(session, 0)
    observer.sample("initial", "initial")
    title = lambda text: lambda frame, context: text in context["windows"]
    step = observer.action
    if app == "xterm":
        step("type_token", lambda: driver.text(goal["token"]))
        step("submit", lambda: driver.key("Return"), title("AI XTERM SAVED"))
    elif app == "chromium":
        step("focus_address", lambda: driver.chord("Control_L", "l"))
        step("type_url", lambda: driver.text(goal["url"]))
        step("navigate", lambda: driver.key("Return"), title("AI FORM READY"))
        step("type_token", lambda: driver.text(goal["token"]))
        step("submit", lambda: driver.key("Return"), title("AI FORM SAVED"))
    elif app == "calc":
        step("type_a", lambda: driver.text(str(goal["a"])))
        step("next_row", lambda: driver.key("Return"))
        step("type_b", lambda: driver.text(str(goal["b"])))
        step("commit_b", lambda: driver.key("Return"))
        step("save", lambda: driver.chord("Control_L", "s"), title("Confirm File Format"))
        step("confirm_format", lambda: driver.key("Return"),
             lambda frame, context: "Confirm File Format" not in context["windows"])
    else:
        # Readiness is a public image condition; no SVG-file read in this function.
        step("selection_tool", lambda: driver.key("F1"), lambda frame, context: red_bbox(frame) is not None)
        step("select_all", lambda: driver.chord("Control_L", "a"),
             lambda frame, context: selection_visible(frame))
        bbox = red_bbox(observer.last)
        if bbox is None:
            raise RuntimeError("Visible target missing")
        x, y = (bbox[0] + bbox[2]) // 2, (bbox[1] + bbox[3]) // 2

        def moved(frame, context):
            current = red_bbox(frame)
            return current is not None and current[0] >= bbox[0] + 2

        step("drag_right", lambda: driver.drag(x, y, x + goal["dx"], y), moved)
        step("save", lambda: driver.chord("Control_L", "s"))
    return {"input_events": driver.input_events, "logical_ops": driver.logical_ops}


def evaluate(app, output, goal):
    """Post-controller scoring only; never used to choose or retry an input."""
    deadline = time.perf_counter() + PARAMETERS["oracle_timeout_s"]
    actual = None
    while True:
        try:
            if app == "xterm":
                actual = output.read_text()
                ok = actual == goal["token"]
            elif app == "chromium":
                actual = urllib.parse.parse_qs(output.read_text())
                ok = actual == {"value": [goal["token"]]}
            elif app == "calc":
                wb = load_workbook(output, read_only=True, data_only=False)
                actual = [wb.active["A1"].value, wb.active["A2"].value]
                wb.close()
                ok = actual == [goal["a"], goal["b"]]
            else:
                root = ET.parse(output).getroot()
                rects = root.findall("{http://www.w3.org/2000/svg}rect")
                if len(rects) != 1:
                    raise ValueError("Expected exactly one rectangle")
                rect = rects[0]
                actual = {k: rect.attrib.get(k) for k in ("x", "y", "width", "height", "transform")}
                # Task contract: move right and preserve y and size, not exact motor gain.
                ok = (actual["transform"] is None and float(actual["x"]) > 50.5
                      and abs(float(actual["y"]) - 50) < .1
                      and abs(float(actual["width"]) - 40) < .1
                      and abs(float(actual["height"]) - 30) < .1)
            if ok:
                return {"success": True, "actual": actual, "known_ns": time.perf_counter_ns()}
        except Exception as exc:
            actual = {"error": repr(exc)}
        if time.perf_counter() >= deadline:
            return {"success": False, "actual": actual, "known_ns": time.perf_counter_ns()}
        time.sleep(.01)


def audit_and_archive(observer, out):
    frames_dir = out / "frames"
    frames_dir.mkdir()
    records = []
    false_suppressions = reconstruction_errors = missed_changes = 0
    prior = None
    eligible = 0
    for sample in observer.samples:
        source, received, update = sample["source"], sample["delivered"], sample["update"]
        # Independent numpy full-array equality, after controller/timing completed.
        same_shape = prior is not None and (prior.width, prior.height, prior.mode) == (source.width, source.height, source.mode)
        equal_prior = same_shape and np.array_equal(np.frombuffer(prior.pixels, dtype=np.uint8),
                                                   np.frombuffer(source.pixels, dtype=np.uint8))
        restored = ((source.width, source.height, source.mode) == (received.width, received.height, received.mode)
                    and np.array_equal(np.frombuffer(source.pixels, dtype=np.uint8),
                                       np.frombuffer(received.pixels, dtype=np.uint8)))
        suppressed = update.frame is None
        eligible += int(equal_prior)
        false_suppressions += int(suppressed and not equal_prior)
        missed_changes += int(suppressed and not restored)
        reconstruction_errors += int(not restored)
        digest = hashlib.sha256(f"{source.width},{source.height},{source.mode}:".encode() + source.pixels).hexdigest()
        png = frames_dir / f"{digest}.png"
        if not png.exists():
            image_for(source).save(png)
        record = {k: v for k, v in sample.items() if k not in ("source", "delivered", "update")}
        record.update(sequence=update.sequence, base_sequence=update.base_sequence,
                      action_id=update.action_id, observed_ns=update.observed_ns,
                      context=update.context, compare_ns=update.compare_ns,
                      suppressed=suppressed, exact_unchanged=bool(equal_prior),
                      reconstructed_exactly=bool(restored), sha256=digest,
                      width=source.width, height=source.height, mode=source.mode)
        records.append(record)
        prior = source
    (out / "observations.jsonl").write_text("".join(json.dumps(r, sort_keys=True) + "\n" for r in records))
    # Derive changed feedback after timing. Do not add a comparison to O0's
    # measured execution or confuse fresh-but-unchanged receipt with an effect.
    for action in observer.actions:
        positions = [i for i, r in enumerate(records) if r["action_id"] == action["action_id"]]
        before = records[positions[0] - 1] if positions and positions[0] else None
        changed = next((records[i] for i in positions if before is not None
                        and (records[i]["sha256"] != before["sha256"]
                             or records[i]["context"] != before["context"])), None)
        action["first_changed_feedback_ms"] = (None if changed is None else
                                                (changed["ready_ns"] - action["issued_ns"]) / 1e6)
    write_json(out / "actions.json", observer.actions)
    count = len(records)
    forwarded = [r for r in records if not r["suppressed"]]
    return {"candidate_observations": count, "model_visible_observations": len(forwarded),
            "suppressed_observations": count - len(forwarded),
            "captured_pixels": sum(r["width"] * r["height"] for r in records),
            "model_visible_pixels": sum(r["width"] * r["height"] for r in forwarded),
            "exact_repeat_opportunities": eligible,
            "same_trace_o1_visible": count - eligible,
            "false_suppressions": false_suppressions, "missed_changes": missed_changes,
            "reconstruction_errors": reconstruction_errors,
            "hash_compute_ns_in_gate": 0, "compare_ns_total": sum(r["compare_ns"] for r in records)}


def episode(app, seed, strategy, chromium, out):
    out.mkdir(parents=True, exist_ok=False)
    session = None
    server = None
    observer = None
    output = None
    goal = None
    start = time.perf_counter_ns()
    result = {"app": app, "seed": seed, "strategy": strategy, "success": False,
              "model_calls": 0, "parameters": PARAMETERS}
    try:
        session = Session()
        goal, output, server = prepare(session, app, seed, chromium)
        result["goal"] = goal
        result["launch_ms"] = (time.perf_counter_ns() - start) / 1e6
        observer = Observer(session, strategy, f"{app}-{seed}-{strategy}")
        task_start = time.perf_counter_ns()
        result.update(controller(session, app, goal, observer))
        returned = time.perf_counter_ns()
        score = evaluate(app, output, goal)
        result.update(success=score["success"], oracle=score,
                      controller_wall_ms=(returned - task_start) / 1e6,
                      task_wall_ms=(score["known_ns"] - task_start) / 1e6,
                      oracle_after_controller_ms=(score["known_ns"] - returned) / 1e6)
    except Exception:
        result["error"] = traceback.format_exc()
        result["elapsed_to_error_ms"] = (time.perf_counter_ns() - start) / 1e6
        if output is not None and goal is not None:
            result["oracle_after_controller_error"] = evaluate(app, output, goal)
    finally:
        if output is not None and output.exists():
            shutil.copy2(output, out / output.name)
        if observer is not None:
            result.update(audit_and_archive(observer, out))
        if server is not None:
            server.shutdown()
            server.server_close()
        if session is not None:
            log = session.tmp / "application.log"
            if log.exists():
                shutil.copy2(log, out / "application.txt")
            session.close()
            # Only this episode's mkdtemp directory, after its process groups stop.
            shutil.rmtree(session.tmp)
    result["audit_pass"] = all(result.get(k) == 0 for k in
                                ("false_suppressions", "missed_changes", "reconstruction_errors"))
    write_json(out / "result.json", result)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", action="store_true")
    parser.add_argument("--freeze", action="store_true")
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--phase", choices=("development", "fresh"), default="development")
    parser.add_argument("--apps", nargs="+", choices=APPS, default=list(APPS))
    parser.add_argument("--strategies", nargs="+", choices=("O0", "O1"), default=["O0", "O1"])
    parser.add_argument("--pairs", type=int, default=2)
    parser.add_argument("--seed", type=int, default=1101)
    parser.add_argument("--chromium", default=shutil.which("chromium"))
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if not args.chromium:
        parser.error("--chromium must point to a headful Chromium-family binary")
    args.out = args.out.resolve()
    if args.worker:
        def interrupted(signum, stack):
            raise TimeoutError("Worker interrupted; preserving evidence and closing its X session")
        signal.signal(signal.SIGTERM, interrupted)
        result = episode(args.apps[0], args.seed, args.strategies[0], args.chromium, args.out)
        print(json.dumps({k: result.get(k) for k in ("app", "seed", "strategy", "success", "audit_pass",
                                                   "candidate_observations", "suppressed_observations", "error")}))
        return
    if args.freeze:
        args.out.mkdir(parents=True, exist_ok=False)
        write_json(args.out / "freeze.json", environment(args.chromium))
        snapshots = args.out / "source"
        snapshots.mkdir()
        for name in ("exact_gate.py", "gui_suite.py", "test_exact_gate.py", "analyze.py", "PROTOCOL.md"):
            shutil.copy2(HERE / name, snapshots / name)
        shutil.copy2(HERE.parent / "real_apps_v1" / "real_app_suite_v1.py", snapshots / "real_app_suite_v1.py")
        print(str(args.out / "freeze.json"))
        return
    if args.phase == "fresh":
        if args.manifest is None:
            parser.error("Fresh evaluation requires a pre-existing freeze manifest")
        frozen = json.loads(args.manifest.read_text())
        if frozen["code_sha256"] != code_hashes() or frozen["parameters"] != PARAMETERS:
            parser.error("Frozen code or parameters changed; new development/freeze required")
    current_environment = environment(args.chromium)
    if args.phase == "fresh" and frozen["versions"] != current_environment["versions"]:
        parser.error("App/dependency versions changed after freeze")
    args.out.mkdir(parents=True, exist_ok=False)
    write_json(args.out / "environment.json", current_environment)
    schedule = [(app, args.seed + i) for i in range(args.pairs) for app in args.apps]
    random.Random(args.seed).shuffle(schedule)
    write_json(args.out / "schedule.json", {"phase": args.phase, "pairs_per_app": args.pairs,
                                           "strategies": args.strategies, "pairs": schedule,
                                           "manifest": str(args.manifest) if args.manifest else None})
    results = []
    for index, (app, seed) in enumerate(schedule):
        order = list(args.strategies)
        if args.phase == "fresh" and (seed - args.seed) % 2:
            order.reverse()
        for strategy in order:
            path = args.out / f"{app}-{seed}-{strategy}"
            cmd = [sys.executable, str(Path(__file__).resolve()), "--worker", "--apps", app,
                   "--strategies", strategy, "--seed", str(seed), "--chromium", args.chromium,
                   "--out", str(path)]
            try:
                process = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                try:
                    worker_stdout, worker_stderr = process.communicate(timeout=90)
                except subprocess.TimeoutExpired:
                    process.terminate()
                    try:
                        worker_stdout, worker_stderr = process.communicate(timeout=15)
                    except subprocess.TimeoutExpired:
                        process.kill()
                        worker_stdout, worker_stderr = process.communicate()
                    raise TimeoutError("Worker deadline exceeded: " + worker_stderr)
                result = json.loads((path / "result.json").read_text())
                if process.returncode:
                    result.update(success=False, worker_stderr=worker_stderr)
            except Exception:
                result = {"app": app, "seed": seed, "strategy": strategy, "success": False,
                          "error": traceback.format_exc()}
            result["pair_index"] = index
            result["relative_path"] = path.name
            results.append(result)
            with (args.out / "runs.jsonl").open("a") as log:
                log.write(json.dumps(result, sort_keys=True) + "\n")
            print(json.dumps({k: result.get(k) for k in ("app", "seed", "strategy", "success", "audit_pass",
                                                         "candidate_observations", "suppressed_observations", "error")}), flush=True)
            if not result.get("success") or not result.get("audit_pass"):
                write_json(args.out / "STOPPED.json", {"reason": "correctness/audit gate; diagnose before comparisons",
                                                       "failed": path.name})
                raise SystemExit(2)
    write_json(args.out / "COMPLETE.json", {"runs": len(results), "success": sum(r["success"] for r in results)})


if __name__ == "__main__":
    main()
