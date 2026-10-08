import hashlib,io,json,statistics,tarfile
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).parent
manifest=json.loads((root/'manifest.json').read_text())
raw=(root/'raw.tar.gz').read_bytes()
assert hashlib.sha256(raw).hexdigest()==manifest['archive_sha256']
with tarfile.open(fileobj=io.BytesIO(raw),mode='r:gz') as t:
 files={m.name:t.extractfile(m).read() for m in t.getmembers() if m.isfile()}
assert set(files)==set(manifest['files'])
for name,data in files.items(): assert hashlib.sha256(data).hexdigest()==manifest['files'][name],name
def read(folder,name): return json.loads(files[f'{folder}/{name}'].decode('utf-8-sig'))
for folder in ['calc-runtime-pacing-02','calc-public-pacing-01']:
 plan=read(folder,'plan.json'); rows=read(folder,'rows.json'); result=read(folder,'result.json')
 assert len(rows)==len(plan['cases'])==32
 book=load_workbook(io.BytesIO(files[f'{folder}/sheet.xlsx']),data_only=False)
 for i,row in enumerate(rows,1):
  assert row['row']==i and {k:row[k] for k in ('text','gap_ms')}==plan['cases'][i-1]
  assert row['actual']==str(book.active.cell(i,1).value)
  assert row['correct']==(row['text']==row['actual'])
  execution=row['execution']
  assert all(r['verified'] and not r['keys_down'] and not r['buttons_down'] for r in execution['releases'])
  assert execution['releases']
  assert row['elapsed_ms']==(execution['ended_ns']-execution['started_ns'])/1e6
  if folder=='calc-public-pacing-01':
   call=read(folder,f'call-{i}.json')
   assert call['response']['result']['status']=='completed'
   assert call['response']['result']['execution']==execution
   assert call['request']['ops'][1]=={'op':'text','text':row['text'],'gap_ms':row['gap_ms']}
 for aggregate in result['summary']:
  arm=[r for r in rows if r['gap_ms']==aggregate['gap_ms']]
  assert aggregate['n']==len(arm)==8
  assert aggregate['correct']==sum(r['correct'] for r in arm)
  assert aggregate['median_execution_ms']==statistics.median(r['elapsed_ms'] for r in arm)
 assert read(folder,'cleanup.json')['processes_stopped'] is True
 book.close()
assert read('calc-runtime-pacing-01','failure.json')['status']=='SETUP_FAILED'
print(f'PASS {len(files)} retained files; 64 saved-value checks, 32 public receipts; failures retained')
