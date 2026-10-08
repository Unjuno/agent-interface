import hashlib,json,os,signal,socket,subprocess,sys,time
from pathlib import Path
from Xlib import display,X
from openpyxl import Workbook,load_workbook
archive,bundle,case=map(lambda p:Path(p).resolve(),sys.argv[1:])
case.mkdir(exist_ok=False)
sys.path.insert(0,str(archive))
from runtime.guarded_x11_v1.bridge import read_window_title
def save(name,row):(case/name).write_text(json.dumps(row,indent=2)+'\n')
children=[];host_exit=None;connection=None
save('allocation.json',{'seed':1001077,'pid':os.getpid(),'source':'49c965df4','archive_sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'retry_budget':0,'expected':[317,529]})
for n in range(27000,27100):
    if Path(f'/tmp/.X11-unix/X{n}').exists() or Path(f'/tmp/.X{n}-lock').exists():continue
    sock=socket.socket(socket.AF_UNIX)
    try:sock.connect('\0'+f'/tmp/.X11-unix/X{n}')
    except OSError:break
    finally:sock.close()
else:raise RuntimeError('no unused display')
env={k:v for k,v in os.environ.items() if k not in ('WAYLAND_DISPLAY','DBUS_SESSION_BUS_ADDRESS')}
env.update(DISPLAY=f':{n}',GDK_BACKEND='x11',QT_QPA_PLATFORM='xcb',LANG='C.UTF-8',LC_ALL='C.UTF-8',SAL_USE_VCLPLUGIN='gen')
for key,folder in [('HOME','home'),('XDG_CONFIG_HOME','config'),('XDG_CACHE_HOME','cache'),('XDG_DATA_HOME','data'),('XDG_RUNTIME_DIR','run')]:
    path=case/folder;path.mkdir(mode=0o700);env[key]=str(path)
def launch(label,args,stdin=subprocess.DEVNULL):
    p=subprocess.Popen(args,env=env,stdin=stdin,stdout=(case/(label+'.stdout')).open('xb'),stderr=(case/(label+'.stderr')).open('xb'),start_new_session=True)
    children.append((label,p));return p
try:
    xvfb=launch('xvfb',['Xvfb',f':{n}','-screen','0','1280x800x24','-ac','-nolisten','tcp'])
    deadline=time.monotonic()+5
    while True:
        if xvfb.poll() is not None:raise RuntimeError('Xvfb exited')
        try:connection=display.Display(f':{n}');break
        except Exception:
            if time.monotonic()>deadline:raise
            time.sleep(.02)
    wm=launch('openbox',['openbox','--config-file','/etc/xdg/openbox/rc.xml'])
    atom=connection.intern_atom('_NET_SUPPORTING_WM_CHECK')
    deadline=time.monotonic()+5
    while connection.screen().root.get_full_property(atom,X.AnyPropertyType) is None:
        if wm.poll() is not None or time.monotonic()>deadline:raise RuntimeError('WM not ready')
        time.sleep(.02)
    workbook=case/'sheet-1001077.xlsx';Workbook().save(workbook)
    save('original-workbook.json',{'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    profile=case/'lo-profile';(profile/'user').mkdir(parents=True)
    (profile/'user'/'registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="ShowTipOfTheDay" oor:op="fuse"><value>false</value></prop></item></oor:items>')
    app=launch('calc',['libreoffice','--norestore','--nodefault','--nolockcheck',f'-env:UserInstallation={profile.as_uri()}','--calc',str(workbook)])
    deadline=time.monotonic()+30;target=None
    while target is None:
        prop=connection.screen().root.get_full_property(connection.intern_atom('_NET_CLIENT_LIST'),X.AnyPropertyType)
        for wid in ([] if prop is None else prop.value):
            win=connection.create_resource_object('window',int(wid))
            legacy=win.get_wm_name()
            title=read_window_title(connection,win) or ''
            with (case/'window-samples.jsonl').open('a') as sample:
                sample.write(json.dumps({'ns':time.monotonic_ns(),'window':int(wid),'legacy_title':str(legacy) if legacy is not None else None,'portable_title':title})+'\n')
            if isinstance(title,bytes):title=title.decode('utf-8','replace')
            if workbook.name in title:target=int(wid);break
        if time.monotonic()>deadline:raise RuntimeError('Calc window not ready')
        time.sleep(.05)
    subprocess.run(['wmctrl','-ia',hex(target)],env=env,check=True)
    save('targets.json',{'app':target})
    config={'host':{'command':sys.executable,'args':[str(archive),'relay','--','--targets',str(case/'targets.json'),'--output-directory',str(case/'calls'),'--display',f':{n}','--session-mode','persistent-x11'],'evidenceDirectory':str(case/'host')},'route':'direct-post','exchangeDirectory':str(case/'exchange'),'primaryOptions':{'observationArguments':{'target':'app','frame':'screen_physical_px','region':[0,0,1280,800],'compact':True,'report_refs':True}}}
    save('cli-config.json',config)
    node=subprocess.Popen(['/home/taka/.volta/bin/node',str(bundle/'primary_stdio.mjs'),'--config',str(case/'cli-config.json')],env=env,stdout=(case/'primary-stream.jsonl').open('xb'),stderr=(case/'primary-stream.stderr').open('xb'))
    children.append(('primary',node))
    save('owner.json',{'pid':os.getpid(),'display':f':{n}','window':target,'children':{name:p.pid for name,p in children},'started_ns':time.monotonic_ns()})
    print('READY OWNER / original stdin / raw primary stream',flush=True)
    host_exit=node.wait(timeout=900)
    if host_exit:raise RuntimeError('primary exited '+str(host_exit))
except Exception as error:
    save('exception.json',{'error':repr(error),'replay_allowed':False});raise
finally:
    if connection is not None:connection.close()
    codes=[]
    for name,p in reversed(children):
        if p.poll() is None:
            if name=='primary':p.terminate()
            else:os.killpg(p.pid,signal.SIGTERM)
        try:code=p.wait(timeout=5)
        except subprocess.TimeoutExpired:
            if name=='primary':p.kill()
            else:os.killpg(p.pid,signal.SIGKILL)
            code=p.wait(timeout=5)
        codes.append({'name':name,'pid':p.pid,'returncode':code})
    save('cleanup.json',{'host_exit':host_exit,'children':codes,'ended_ns':time.monotonic_ns()})
    # No application/file oracle is available to the controller before terminal.
    if 'workbook' in locals() and workbook.exists():
        book=load_workbook(workbook,read_only=True,data_only=True)
        cells={cell.coordinate:cell.value for row in book.active for cell in row if cell.value is not None}
        book.close();save('evaluation.json',{'expected':{'A1':317,'A2':529},'actual_nonempty_cells':cells,'success':cells=={'A1':317,'A2':529},'after_all_owned_processes_terminal':True,'workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    print('TERMINAL OWNER',flush=True)

