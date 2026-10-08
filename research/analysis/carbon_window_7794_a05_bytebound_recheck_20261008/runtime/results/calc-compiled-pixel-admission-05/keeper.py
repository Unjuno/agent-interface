import hashlib,json,os,signal,socket,subprocess,sys,time
from pathlib import Path
from Xlib import display,X,error
from openpyxl import Workbook,load_workbook
archive,case=map(lambda p:Path(p).resolve(),sys.argv[1:3])
seed=int(sys.argv[3]);reuse=False
case.mkdir(exist_ok=False)
sys.path.insert(0,str(archive))
from runtime.guarded_x11_v1.bridge import read_window_title
from runtime.guarded_x11_v1.handles import ALIAS
NATIVE_ALIASES=('sheet_context','format_context')
if any(ALIAS.fullmatch(alias) is None for alias in NATIVE_ALIASES):
    raise ValueError('invalid native aliases before display allocation')
def save(name,row):(case/name).write_text(json.dumps(row,indent=2)+'\n')
children=[];host_exit=None;connection=None;bridge=None
node_binary=Path('/home/taka/.volta/tools/image/node/24.13.1/bin/node')
node_preflight=subprocess.run([str(node_binary),'--version'],env={**os.environ,'HOME':str(case)},capture_output=True,text=True,check=True)
save('node-preflight.json',{'binary':str(node_binary),'sha256':hashlib.sha256(node_binary.read_bytes()).hexdigest(),'version':node_preflight.stdout.strip(),'returncode':node_preflight.returncode,'isolated_home':str(case)})
save('allocation.json',{'seed':seed,'reuseReviewedImages':reuse,'pid':os.getpid(),'source':json.loads((archive.parent/'MANIFEST.json').read_text())['source_revision'],'archive_sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'retry_budget':0,'expected':[317,529]})
for n in range(27000,27100):
    if Path(f'/tmp/.X11-unix/X{n}').exists() or Path(f'/tmp/.X{n}-lock').exists():continue
    sock=socket.socket(socket.AF_UNIX)
    try:sock.connect('\0'+f'/tmp/.X11-unix/X{n}')
    except OSError:break
    finally:sock.close()
else:raise RuntimeError('no unused display')
env={k:v for k,v in os.environ.items() if k not in ('WAYLAND_DISPLAY','DBUS_SESSION_BUS_ADDRESS')}
env.update(DISPLAY=f':{n}',GDK_BACKEND='x11',QT_QPA_PLATFORM='xcb',LANG='C.UTF-8',LC_ALL='C.UTF-8',SAL_USE_VCLPLUGIN='gen')
for key,folder in [('HOME','home'),('XDG_CONFIG_HOME','config'),('XDG_CACHE_HOME','cache'),('XDG_DATA_HOME','data'),('XDG_RUNTIME_DIR','run')]:
    path=case/folder;path.mkdir(mode=0o700);env[key]=str(path)
def launch(label,args,stdin=subprocess.DEVNULL):
    p=subprocess.Popen(args,env=env,stdin=stdin,stdout=(case/(label+'.stdout')).open('xb'),stderr=(case/(label+'.stderr')).open('xb'),start_new_session=True)
    children.append((label,p));return p
try:
    xvfb=launch('xvfb',['Xvfb',f':{n}','-screen','0','1280x800x24','-ac','-nolisten','tcp'])
    deadline=time.monotonic()+5
    while True:
        if xvfb.poll() is not None:raise RuntimeError('Xvfb exited')
        try:connection=display.Display(f':{n}');break
        except Exception:
            if time.monotonic()>deadline:raise
            time.sleep(.02)
    wm=launch('openbox',['openbox','--config-file','/etc/xdg/openbox/rc.xml'])
    atom=connection.intern_atom('_NET_SUPPORTING_WM_CHECK')
    deadline=time.monotonic()+5
    while connection.screen().root.get_full_property(atom,X.AnyPropertyType) is None:
        if wm.poll() is not None or time.monotonic()>deadline:raise RuntimeError('WM not ready')
        time.sleep(.02)
    workbook=case/f'sheet-{seed}.xlsx';Workbook().save(workbook)
    save('original-workbook.json',{'sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    profile=case/'lo-profile';(profile/'user').mkdir(parents=True)
    (profile/'user'/'registrymodifications.xcu').write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="ShowTipOfTheDay" oor:op="fuse"><value>false</value></prop></item></oor:items>')
    app=launch('calc',['libreoffice','--norestore','--nodefault','--nolockcheck',f'-env:UserInstallation={profile.as_uri()}','--calc',str(workbook)])
    deadline=time.monotonic()+30;target=None
    while target is None:
        prop=connection.screen().root.get_full_property(connection.intern_atom('_NET_CLIENT_LIST'),X.AnyPropertyType)
        for wid in ([] if prop is None else prop.value):
            win=connection.create_resource_object('window',int(wid))
            try:
                legacy=win.get_wm_name()
                title=read_window_title(connection,win) or ''
            except error.BadWindow as stale_error:
                with (case/'window-discovery-refusals.jsonl').open('a') as sample:
                    sample.write(json.dumps({'ns':time.monotonic_ns(),'window':int(wid),'error':repr(stale_error),'authority_granted':False})+'\n')
                continue
            with (case/'window-samples.jsonl').open('a') as sample:
                sample.write(json.dumps({'ns':time.monotonic_ns(),'window':int(wid),'legacy_title':str(legacy) if legacy is not None else None,'portable_title':title})+'\n')
            if isinstance(title,bytes):title=title.decode('utf-8','replace')
            if workbook.name in title:target=int(wid);break
        if time.monotonic()>deadline:raise RuntimeError('Calc window not ready')
        time.sleep(.05)
    subprocess.run(['wmctrl','-ia',hex(target)],env=env,check=True)
    save('targets.json',{'app':target})
    from runtime.guarded_x11_v1.bridge import NativeHandleBridge
    bridge=NativeHandleBridge(f':{n}',{'app':target},'app',case/'bridge')
    # Read-only actual backend preflight, not input authority or a lease.
    before=bridge.backend.emissions
    invalid={'ops':[{'op':'focus','target':'app'},{'op':'key_chord','keys':['CTRL','HOME']},{'op':'release_all'}]}
    try:
        bridge.backend.preflight(invalid)
    except Exception as error:
        if 'unmapped key HOME' not in str(error):raise
        negative={'status':'refused','detail':str(error)}
    else:raise RuntimeError('negative HOME preflight unexpectedly accepted')
    planned={'ops':[{'op':'focus','target':'app'},{'op':'key_chord','keys':['CTRL','Home']},
        {'op':'text','text':'317'},{'op':'key_chord','keys':['ENTER']},
        {'op':'text','text':'529'},{'op':'key_chord','keys':['ENTER']},
        {'op':'key_chord','keys':['CTRL','s']},{'op':'key_chord','keys':['ENTER']},
        {'op':'release_all'}]}
    mapping=bridge.backend.preflight(planned)
    if before!=bridge.backend.emissions:raise RuntimeError('preflight emitted input')
    save('backend-preflight.json',{'valid_status':'accepted','negative':negative,
        'planned':planned,'keyboard_mapping_sha256':hashlib.sha256(repr(mapping).encode()).hexdigest(),
        'emissions_before':before,'emissions_after':bridge.backend.emissions,
        'authority_granted':False,'scope':'backend binding feasibility only; actual dispatch repeats admission/mapping checks'})
    save('owner.json',{'pid':os.getpid(),'display':f':{n}','window':target,'children':{name:p.pid for name,p in children},'started_ns':time.monotonic_ns()})
    (case/'commands').mkdir();(case/'replies').mkdir()
    print('READY / existing command-file allocation protocol',flush=True)
    index=1;source=None;image=None;refs={};run_used=False;confirm_used=False
    stop_at=time.monotonic()+900
    while True:
        if time.monotonic()>stop_at:raise RuntimeError('bounded owner lifetime exhausted')
        if any(p.poll() is not None for name,p in children):raise RuntimeError('owned child exited')
        path=case/'commands'/f'{index:03d}.json'
        if not path.exists():time.sleep(.02);continue
        request=json.loads(path.read_text());started=time.monotonic_ns();op=request['op']
        if op=='observe':
            source=bridge.observe();image=bridge.history[source['sequence']][1].copy()
            reply={'observation':source,'focused_candidate':bridge.focused_client_window()}
        elif op=='mint':
            if source is None:raise RuntimeError('primary reviewed source required')
            alias=request['alias'];point=request['point']
            if alias not in NATIVE_ALIASES:raise ValueError('bounded aliases only')
            minted=bridge.mint_reference(alias,source['sequence'],point,region_size=(24,24))
            box=[point[0]-12,point[1]-12,point[0]+12,point[1]+12]
            refs[alias]={'offset':minted['offset'],'box':box,'pixels':image.crop(box).tobytes()}
            reply=minted
        elif op=='run':
            if run_used:raise RuntimeError('single graph allocation; no retry')
            run_used=True
            sys.path.insert(0,str(Path(__file__).parent))
            from methods import run
            reply=run(bridge,refs,request['regions'])
        elif op=='review_window':
            row=bridge.review_window(request['window_id']);reply=row;refs={}
            if row['status']=='reviewed':
                source=row['observation'];image=bridge.history[source['sequence']][1].copy()
        elif op=='confirm':
            if confirm_used or 'format_context' not in refs:raise RuntimeError('one explicit confirmation with freshly grounded modal')
            confirm_used=True
            reply=bridge.keyboard('format_context',refs['format_context']['offset'],
                tail=[{'op':'key_chord','keys':['ENTER']}],
                expires_at_ns=time.monotonic_ns()+2_000_000_000)
        elif op=='close':reply={'status':'closing','replay_allowed':False}
        else:raise ValueError('unsupported bounded operation')
        save(f'replies/{index:03d}.json',{'request':request,'started_ns':started,'ended_ns':time.monotonic_ns(),'reply':reply})
        print('REPLY '+str(index)+' '+op,flush=True);index+=1
        if op=='close':break
    bridge.close();bridge=None;host_exit=0

except Exception as error:
    save('exception.json',{'error':repr(error),'replay_allowed':False});raise
finally:
    if bridge is not None:bridge.close()
    if connection is not None:connection.close()
    codes=[]
    for name,p in reversed(children):
        if p.poll() is None:
            if name=='primary':p.terminate()
            else:os.killpg(p.pid,signal.SIGTERM)
        try:code=p.wait(timeout=5)
        except subprocess.TimeoutExpired:
            if name=='primary':p.kill()
            else:os.killpg(p.pid,signal.SIGKILL)
            code=p.wait(timeout=5)
        codes.append({'name':name,'pid':p.pid,'returncode':code})
    save('cleanup.json',{'host_exit':host_exit,'children':codes,'ended_ns':time.monotonic_ns()})
    # No application/file oracle is available to the controller before terminal.
    if 'workbook' in locals() and workbook.exists():
        book=load_workbook(workbook,read_only=True,data_only=True)
        cells={cell.coordinate:cell.value for row in book.active for cell in row if cell.value is not None}
        book.close();save('evaluation.json',{'expected':{'A1':317,'A2':529},'actual_nonempty_cells':cells,'success':cells=={'A1':317,'A2':529},'after_all_owned_processes_terminal':True,'workbook_sha256':hashlib.sha256(workbook.read_bytes()).hexdigest()})
    print('TERMINAL OWNER',flush=True)



