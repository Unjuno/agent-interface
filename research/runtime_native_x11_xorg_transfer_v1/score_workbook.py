#!/usr/bin/env python3
import argparse, hashlib, json
from openpyxl import load_workbook
CORPUS=["office","coffee","bookkeeper","committee","parallel","address","success","letterpress","mississippi","assessment","committee","bookkeeping","ffffffffff","aaaaaaaaaa","tttttttttt","ssssssssss"]
p=argparse.ArgumentParser();p.add_argument('--workbook',required=True);p.add_argument('--out',required=True);a=p.parse_args()
wb=load_workbook(a.workbook,data_only=False);ws=wb.active
actual=[ws.cell(i+1,1).value for i in range(len(CORPUS))]
exact=sum(x==y for x,y in zip(actual,CORPUS))
raw=open(a.workbook,'rb').read()
r={'schema':'agent-interface/native-x11-xorg-transfer-score-v1','expected':CORPUS,'actual':actual,'exact_count':exact,'passed':exact==len(CORPUS),'xlsx_sha256':hashlib.sha256(raw).hexdigest(),'xlsx_bytes':len(raw)}
open(a.out,'w').write(json.dumps(r,indent=2)+'\n');print(json.dumps(r,indent=2));raise SystemExit(0 if r['passed'] else 1)
