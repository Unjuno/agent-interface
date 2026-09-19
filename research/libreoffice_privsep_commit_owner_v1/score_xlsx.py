#!/usr/bin/env python3
import argparse, json, hashlib
from pathlib import Path
from openpyxl import load_workbook
p=argparse.ArgumentParser(); p.add_argument('--xlsx',type=Path,required=True); p.add_argument('--out',type=Path,required=True); a=p.parse_args()
b=a.xlsx.read_bytes(); wb=load_workbook(a.xlsx,data_only=False,read_only=True); ws=wb.active
cells={k:ws[k].value for k in ['A1','A2','A3','B1']}; wb.close()
r={'path':str(a.xlsx),'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b),'cells':cells}
a.out.write_text(json.dumps(r,indent=2,sort_keys=True)+'\n'); print(json.dumps(r,sort_keys=True))
