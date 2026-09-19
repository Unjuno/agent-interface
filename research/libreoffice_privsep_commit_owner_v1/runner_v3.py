#!/usr/bin/env python3
from pathlib import Path
import argparse, os, subprocess, time, json, hashlib, shutil, errno, pwd
from openpyxl import Workbook
from Xlib import display
import sys
HERE=Path(__file__).resolve().parent
SOURCE=HERE/'source'
sys.path.insert(0,str(SOURCE))
from office_backend import OfficeX11Backend

EXPECTED_BLOBS={'backend_x11.py':'b4f8e043ce4f8929d446e038418ea0fd3655bab0','office_backend.py':'3aeca10f62fb1366bcdd5fad561509cfcc0d91ab'}

def gitblob(p):
 b=Path(p).read_bytes(); return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def statrow(p, with_hash=True):
 s=Path(p).stat(); r={'uid':s.st_uid,'gid':s.st_gid,'mode':oct(s.st_mode & 0o777),'dev':s.st_dev,'ino':s.st_ino,'size':s.st_size,'mtime_ns':s.st_mtime_ns,'ctime_ns':s.st_ctime_ns}
 if with_hash and Path(p).is_file(): r['sha256']=sha(p)
 return r
def dump(p,x): Path(p).write_text(json.dumps(x,indent=2,sort_keys=True,default=str)+'\n')
def make_xlsx(p,a1,a2,a3=None,b1=None):
 wb=Workbook(); ws=wb.active; ws['A1']=a1; ws['A2']=a2; ws['A3']=a3; ws['B1']=b1; wb.save(p)
def ensure_user(name):
 try: return pwd.getpwnam(name)
 except KeyError:
  subprocess.run(['useradd','-M','-d',f'/tmp/{name}','-s','/bin/bash',name],check=True)
  return pwd.getpwnam(name)
def wait_display(env,name):
 for _ in range(120):
  if subprocess.run(['xdpyinfo','-display',name],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0:return
  time.sleep(.05)
 raise RuntimeError('Xvfb not ready')
def windows(dn):
 d=display.Display(dn); root=d.screen().root; out=[]
 def walk(w,depth=0):
  try: name=w.get_wm_name(); cls=w.get_wm_class(); attrs=w.get_attributes()
  except Exception: name=cls=attrs=None
  if name or cls: out.append({'id':w.id,'name':str(name) if name else None,'class':list(cls) if cls else None,'depth':depth,'map_state':getattr(attrs,'map_state',None) if attrs else None})
  try: kids=w.query_tree().children
  except Exception: kids=[]
  for c in kids: walk(c,depth+1)
 walk(root); d.close(); return out
def find_calc(dn):
 end=time.monotonic()+15
 while time.monotonic()<end:
  for r in windows(dn):
   if r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']): return r['id']
  time.sleep(.05)
 raise RuntimeError('Calc window not found')
def modal_names(rows,calc):
 return sorted(set(r['name'] for r in rows if r['id']!=calc and r['name'] and r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class'])))
def proc_uids(fragment):
 out=[]
 for p in Path('/proc').iterdir():
  if not p.name.isdigit(): continue
  try:
   cmd=(p/'cmdline').read_bytes().replace(b'\0',b' ').decode(errors='replace')
   if fragment not in cmd or 'soffice.bin' not in cmd: continue
   status=(p/'status').read_text(); uid_line=next(x for x in status.splitlines() if x.startswith('Uid:'))
   euid=int(uid_line.split()[2]); out.append({'pid':int(p.name),'euid':euid,'cmd':cmd[:500]})
  except Exception: pass
 return sorted(out,key=lambda x:x['pid'])
def independent_score(xlsx,out):
 sp=subprocess.run([sys.executable,str(HERE/'score_xlsx.py'),'--xlsx',str(xlsx),'--out',str(out)],capture_output=True,text=True)
 if sp.returncode!=0: raise RuntimeError('scorer failed: '+sp.stderr)
 return json.loads(Path(out).read_text())
def publish(staging,target,target_dir):
 tmp=target_dir/f'.publish-{os.getpid()}.tmp'; shutil.copyfile(staging,tmp)
 with tmp.open('rb') as f: os.fsync(f.fileno())
 os.replace(tmp,target)
 fd=os.open(target_dir,os.O_RDONLY); os.fsync(fd); os.close(fd)
def writer_attempt(gui,gui_uid,external,target,result_path):
 code=r'''import os,json,sys,time
src,dst,out=sys.argv[1:]
r={'started_ns':time.monotonic_ns(),'euid':os.geteuid(),'egid':os.getegid(),'src':src,'dst':dst}
try:
 os.replace(src,dst); r.update(ok=True,errno=None,error=None)
except OSError as e:
 r.update(ok=False,errno=e.errno,error=repr(e))
r['ended_ns']=time.monotonic_ns()
open(out,'w').write(json.dumps(r,sort_keys=True)+'\n')
'''
 r=subprocess.run(['runuser','-u',gui,'--',sys.executable,'-c',code,str(external),str(target),str(result_path)],capture_output=True,text=True)
 if r.returncode!=0: raise RuntimeError('writer harness failed: '+r.stderr)
 row=json.loads(Path(result_path).read_text())
 if row['euid']!=gui_uid: raise RuntimeError('writer uid mismatch')
 return row

def run_arm(root,name,dn,adversarial,gui):
 out=root/name; out.mkdir(); pw=ensure_user(gui); uid,gid=pw.pw_uid,pw.pw_gid
 target_dir=out/'authority'; gui_dir=out/'gui'; target_dir.mkdir(); gui_dir.mkdir(); os.chown(gui_dir,uid,gid); os.chmod(gui_dir,0o700); os.chmod(target_dir,0o755)
 target=target_dir/'target.xlsx'; staging=gui_dir/'staging.xlsx'; external=gui_dir/'external.xlsx'
 make_xlsx(target,'seed','old'); make_xlsx(staging,'seed','old'); make_xlsx(external,'external','replacement','external-marker','writer')
 os.chmod(target,0o644)
 for p in [staging,external]: os.chown(p,uid,gid); os.chmod(p,0o600)
 profile=gui_dir/'profile'; home=gui_dir/'home'; xdg=gui_dir/'xdg'
 for p in [profile,home,xdg]: p.mkdir(); os.chown(p,uid,gid); os.chmod(p,0o700)
 plan=statrow(target); result={'arm':name,'adversarial':adversarial,'gui_user':gui,'gui_uid':uid,'gui_gid':gid,'authority_dir_initial':statrow(target_dir,False),'plan_target':plan,'source_blobs':EXPECTED_BLOBS}
 auth=out/'controller.Xauthority'; auth.write_bytes(b'')
 env=os.environ.copy(); env.update(DISPLAY=dn,XAUTHORITY=str(auth)); os.environ['DISPLAY']=dn; os.environ['XAUTHORITY']=str(auth)
 xv=subprocess.Popen(['Xvfb',dn,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(out/'xvfb.stdout').open('w'),stderr=(out/'xvfb.stderr').open('w'))
 ob=lo=None; b=None
 try:
  wait_display(env,dn)
  genv=['env',f'DISPLAY={dn}',f'HOME={home}',f'XDG_RUNTIME_DIR={xdg}',f'USER={gui}',f'LOGNAME={gui}']
  ob=subprocess.Popen(['runuser','-u',gui,'--',*genv,'openbox'],stdout=(out/'openbox.stdout').open('w'),stderr=(out/'openbox.stderr').open('w')); time.sleep(.4)
  profile_url=f'file://{profile}'
  lo=subprocess.Popen(['runuser','-u',gui,'--',*genv,'libreoffice',f'-env:UserInstallation={profile_url}','--nologo','--nodefault','--nofirststartwizard','--norestore',str(staging)],stdout=(out/'libreoffice.stdout').open('w'),stderr=(out/'libreoffice.stderr').open('w'))
  calc=find_calc(dn); result['calc_window_id']=calc; time.sleep(.6); result['soffice_processes']=proc_uids(str(profile))
  if not result['soffice_processes'] or any(p['euid']!=uid for p in result['soffice_processes']): raise RuntimeError('LibreOffice privilege separation not proven')
  b=OfficeX11Backend(dn,{'calc':calc}); b.focus('calc'); b.pointer_move('calc','window_client',80,180); b.pointer_button('left',True); b.pointer_button('left',False); b.key_chord(['CTRL','Home']); b.text('office'); b.key_chord(['ENTER']); b.text('preview'); b.key_chord(['ENTER']); time.sleep(.15)
  b.key_chord(['CTRL','S']); time.sleep(.7); m1=modal_names(windows(dn),calc)
  if m1==['Confirm File Format']: b.key_chord(['ENTER']); result['format_confirm_enter_sent']=True; time.sleep(1.2)
  else: result['format_confirm_enter_sent']=False
  m2=modal_names(windows(dn),calc); result['modal_names_1']=m1; result['modal_names_2']=m2
  result['release']=b.release_all().__dict__; result['backend_emissions']=b.emissions
  result['staging_score']=independent_score(staging,out/'staging_score.json')
  if result['staging_score']['cells']['A1']!='office' or result['staging_score']['cells']['A2']!='preview': raise RuntimeError('staging save incorrect')
  # Final controller check occurs before the adversarial replace attempt.
  checked=statrow(target); result['final_check_target']=checked; result['plan_valid_at_check']=all(checked[k]==plan[k] for k in ['sha256','ino','size'])
  if not result['plan_valid_at_check']: raise RuntimeError('target changed before final check')
  result['writer']=None
  if adversarial:
   wr=writer_attempt(gui,uid,external,target,gui_dir/'writer_result.json'); result['writer']=wr; result['target_after_writer_attempt']=statrow(target); result['external_exists_after_writer_attempt']=external.exists()
  # Revalidate target as frozen in Issue #318 before controller publication.
  current=statrow(target); result['publish_recheck_target']=current; result['publish_recheck_valid']=all(current[k]==plan[k] for k in ['sha256','ino','size'])
  result['publish_started_ns']=time.monotonic_ns()
  if result['publish_recheck_valid']:
   publish(staging,target,target_dir); result['published']=True
  else: result['published']=False
  result['publish_ended_ns']=time.monotonic_ns(); result['final_target']=statrow(target); result['final_score']=independent_score(target,out/'final_score.json'); result['authority_dir_final']=statrow(target_dir,False); result['lo_alive_at_score']=lo.poll() is None
  okcells=result['final_score']['cells']['A1']=='office' and result['final_score']['cells']['A2']=='preview'
  denied=(not adversarial) or (result['writer'] and result['writer']['ok'] is False and result['writer']['errno'] in [errno.EACCES,errno.EPERM])
  unchanged=(not adversarial) or (result['target_after_writer_attempt']['sha256']==plan['sha256'] and result['target_after_writer_attempt']['ino']==plan['ino'])
  result['classification']='PASS_CAPABILITY_COMMIT_OWNER' if result['published'] and okcells and denied and unchanged and result['release']['verified'] else 'FAIL_OR_UNCERTAIN'
  dump(out/'result.json',result); return result
 except Exception as e:
  dump(out/'error.json',{'type':type(e).__name__,'detail':str(e),'monotonic_ns':time.monotonic_ns()}); raise
 finally:
  if b:
   try: b.release_all(); b.close()
   except Exception: pass
  if lo and lo.poll() is None: lo.terminate()
  if lo:
   try: lo.wait(timeout=3)
   except Exception: lo.kill()
  if ob and ob.poll() is None: ob.terminate()
  if ob:
   try: ob.wait(timeout=2)
   except Exception: ob.kill()
  xv.terminate()
  try: xv.wait(timeout=2)
  except Exception: xv.kill()

def main():
 ap=argparse.ArgumentParser(); ap.add_argument('--out',type=Path,required=True); ap.add_argument('--gui-user',default='ai_gui'); args=ap.parse_args()
 for n,g in EXPECTED_BLOBS.items():
  actual=gitblob(SOURCE/n)
  if actual!=g: raise RuntimeError(f'{n} blob mismatch {actual} != {g}')
 if args.out.exists(): raise FileExistsError(args.out)
 args.out.mkdir(parents=True)
 rows=[]
 for name,dn,adv in [('01-stable',':341',False),('02-path-replace-attempt',':342',True)]: rows.append(run_arm(args.out,name,dn,adv,args.gui_user))
 summary={'schema':'agent-interface/libreoffice-privsep-commit-owner-v1','allocation':'c318-privsep-03','order':['stable','path_replace_attempt'],'rows':rows,'source_blobs':EXPECTED_BLOBS,'source_sha256':{p.name:sha(p) for p in [Path(__file__),HERE/'audit_v3.py',HERE/'score_xlsx.py',SOURCE/'backend_x11.py',SOURCE/'office_backend.py'] if p.exists()}}
 dump(args.out/'SUMMARY.json',summary); print(json.dumps({'allocation':summary['allocation'],'classifications':[r['classification'] for r in rows],'writer':rows[1]['writer'],'final_cells':[r['final_score']['cells'] for r in rows]},indent=2))
if __name__=='__main__': main()
