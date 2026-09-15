#!/usr/bin/env python3
"""Private Xvfb/Openbox/LibreOffice orchestration; score only after executor exits."""
from __future__ import annotations
import argparse, hashlib, json, os, subprocess, sys, time
from pathlib import Path
from openpyxl import Workbook
from Xlib import display

HERE = Path(__file__).resolve().parent

def wait_display(env, name):
    for _ in range(80):
        if subprocess.run(['xdpyinfo','-display',name], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode == 0: return
        time.sleep(.05)
    raise RuntimeError('Xvfb did not become ready')

def find_calc_window(display_name: str, timeout_s=8.0) -> int:
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        d = display.Display(display_name)
        def walk(win):
            try: children = win.query_tree().children
            except Exception: return None
            for child in children:
                try: wm_class = child.get_wm_class()
                except Exception: wm_class = None
                if wm_class and any('libreoffice-calc' in str(x).lower() for x in wm_class): return child.id
                nested = walk(child)
                if nested is not None: return nested
            return None
        found = walk(d.screen().root); d.close()
        if found is not None: return found
        time.sleep(.05)
    raise RuntimeError('Calc window not found')

def main() -> int:
    ap=argparse.ArgumentParser(); ap.add_argument('--out', type=Path, required=True); ap.add_argument('--display', default=':202')
    args=ap.parse_args(); args.out.mkdir(parents=True, exist_ok=False)
    xlsx=args.out/'task.xlsx'; wb=Workbook(); ws=wb.active; ws['A1']='seed'; ws['A2']='old'; wb.save(xlsx)
    input_sha=hashlib.sha256(xlsx.read_bytes()).hexdigest()
    auth=args.out/'Xauthority'; auth.write_bytes(b'')
    env=os.environ.copy(); env['DISPLAY']=args.display; env['XAUTHORITY']=str(auth); os.environ['DISPLAY']=args.display; os.environ['XAUTHORITY']=str(auth)
    portable=Path(os.environ['AGENT_INTERFACE_PORTABLE_ORACLE']); x11=Path(os.environ['AGENT_INTERFACE_X11_BACKEND'])
    child_env=env.copy(); child_env['AGENT_INTERFACE_PORTABLE_ORACLE']=str(portable); child_env['AGENT_INTERFACE_X11_BACKEND']=str(x11)
    xvfb=subprocess.Popen(['Xvfb',args.display,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(args.out/'xvfb.stdout').open('w'),stderr=(args.out/'xvfb.stderr').open('w'))
    openbox=lo=None
    try:
        wait_display(env,args.display)
        openbox=subprocess.Popen(['openbox'],env=env,stdout=(args.out/'openbox.stdout').open('w'),stderr=(args.out/'openbox.stderr').open('w'))
        profile=f'file://{args.out}/profile'
        lo=subprocess.Popen(['libreoffice',f'-env:UserInstallation={profile}','--nologo','--nodefault','--nofirststartwizard','--norestore',str(xlsx)],env=env,stdout=(args.out/'libreoffice.stdout').open('w'),stderr=(args.out/'libreoffice.stderr').open('w'))
        window_id=find_calc_window(args.display)
        time.sleep(.4)
        execute_dir=args.out/'execution'
        run=subprocess.run([sys.executable,str(HERE/'experiment.py'),'--display',args.display,'--window-id',str(window_id),'--out',str(execute_dir)],env=child_env,text=True,capture_output=True)
        (args.out/'executor.stdout').write_text(run.stdout,encoding='utf-8'); (args.out/'executor.stderr').write_text(run.stderr,encoding='utf-8'); (args.out/'executor.exitcode').write_text(str(run.returncode)+'\n')
        time.sleep(.8)  # bounded post-save flush; scorer remains causally downstream.
        score=subprocess.run([sys.executable,str(HERE/'score_workbook.py'),'--xlsx',str(xlsx),'--out',str(args.out/'score.json')],text=True,capture_output=True)
        (args.out/'scorer.stdout').write_text(score.stdout,encoding='utf-8'); (args.out/'scorer.stderr').write_text(score.stderr,encoding='utf-8'); (args.out/'scorer.exitcode').write_text(str(score.returncode)+'\n')
        result={
            'schema':'agent-interface/office-x11-calc-run-v0','window_id':window_id,
            'input_xlsx_sha256':input_sha,'executor_exitcode':run.returncode,'scorer_exitcode':score.returncode,
            'source_scope':'private Xvfb/Openbox LibreOffice Calc; no model/provider/native platform claim',
        }
        result['passed']=run.returncode==0 and score.returncode==0
        (args.out/'report.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
        print(json.dumps(result,indent=2)); return 0 if result['passed'] else 1
    finally:
        if lo and lo.poll() is None: lo.terminate()
        if lo:
            try: lo.wait(timeout=3)
            except subprocess.TimeoutExpired: lo.kill()
        if openbox and openbox.poll() is None: openbox.terminate()
        xvfb.terminate()

if __name__ == '__main__': raise SystemExit(main())
