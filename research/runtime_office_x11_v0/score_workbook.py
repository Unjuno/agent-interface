#!/usr/bin/env python3
"""Independent post-execution XLSX scorer. Never imported by the executor."""
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook


def score(path: Path) -> dict:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb.active
    cells = {name: ws[name].value for name in ('A1', 'A2', 'A3')}
    passed = cells == {'A1': 'office', 'A2': 'preview', 'A3': None}
    return {
        'schema': 'agent-interface/office-x11-calc-independent-score-v0',
        'passed': passed,
        'cells': cells,
        'xlsx_sha256': digest,
        'xlsx_bytes': path.stat().st_size,
    }


def main() -> int:
    ap = argparse.ArgumentParser(); ap.add_argument('--xlsx', type=Path, required=True); ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args(); result = score(args.xlsx)
    args.out.write_text(json.dumps(result, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, indent=2)); return 0 if result['passed'] else 1

if __name__ == '__main__': raise SystemExit(main())
