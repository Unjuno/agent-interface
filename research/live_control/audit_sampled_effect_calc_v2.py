"""Replay Calc provenance, shared caller decisions, and independent workbook."""
import hashlib,json,sys
from pathlib import Path
from PIL import Image
from openpyxl import load_workbook
from received_continuation_v1 import start,advance
from append_checkpoint_v1 import load,inspect
from sampled_target_contract_v1 import evaluate
from phased_submit_v1 import execute
H=Path(__file__).resolve().parent;R=H/'results/sampled-effect-calc-02';T=R/'runtime'
sys.path.insert(0,str(H.parent/'observation_tiles'))
from tile_transport import Decoder
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
for n,s in read(R/'plan.json')['sources'].items():assert sha(H/n)==s
for n,s in read(T/'sources.json').items():assert sha(H.parent/n)==s
events=[json.loads(l) for l in (T/'events.jsonl').read_bytes().splitlines()];ep=read(R/'endpoint.json');state=start(ep['socket'])
def replay(r):
 global state
 q,a=r['request'],r['reply'];assert q['after']==state['cursor'] and a['records']==events[q['after']:a['cursor']]
 state=advance(state,ep['socket'],q['after'],a);assert state==r.get('state',{}).get('continuation',r.get('continuation'))
replay(read(R/'initial.json'));calls=read(R/'calls.json')
for c in calls:replay(c['result'])
assert load(R/'journal.jsonl')==calls[-1]['result']['state'];assert inspect(R/'journal.jsonl')[1]==1+2*len(calls)
replay(read(R/'finish.json'));models=[]
for row in read(R/'turns.json'):
 n=row['turn'];d=R/f'model-{n}';plan=read(d/'plan.json');process=read(d/'process.json')
 assert process['exit_code']==0 and plan['image_sha256']==sha(T/Path(row['source']['image']).name)
 assert plan['runner_sha256']==sha(H/'model_context_runner_v1.py') and plan['instructions_sha256']==sha(H/'screenshot_responder_v1.txt')
 assert (d/'prompt.txt').read_text(encoding='utf-8')==(R/f'prompt-{n}.txt').read_text(encoding='utf-8')
 assert json.loads((d/'prompt.txt').read_text(encoding='utf-8').split('Evidence: ')[1])==row['feedback']
 raw=(d/'events.jsonl').read_bytes().splitlines(keepends=True);arr=[json.loads(l) for l in (d/'arrivals.jsonl').read_bytes().splitlines()]
 assert len(raw)==len(arr)==4
 for i,(line,a) in enumerate(zip(raw,arr)):assert a['line']==i and a['sha256']==hashlib.sha256(line).hexdigest() and a['bytes']==len(line)
 e=[json.loads(l) for l in raw];assert json.loads(e[2]['item']['text'])==row['proposal']==read(R/f'proposal-{n}.json')
 models.append({'turn':n,'usage':e[3]['usage'],'runner_s':(process['exited_ns']-process['started_ns'])/1e9})
 if row['proposal']['kind']=='verify':continue
 subset=[c['result'] for c in calls[row['calls_begin']:row['calls_end']]]
 assert row['fresh']==subset[1]['state']['continuation']['observation'] and row['clock']==subset[2]['state']['last_resolution']['clock']
 with Image.open(T/Path(row['source']['image']).name) as old,Image.open(T/Path(row['fresh']['image']).name) as new:
  assert evaluate(row['contract'],{'intent':row['contract']['name'],'execute_once':True},row['source'],row['fresh'],old,new,row['clock']['runtime_ns'])==row['checked']
 responses=iter(subset[3:]);used=[]
 def fake(spec):
  r=next(responses);used.append(r);assert all(r['request']['command'][k]==v for k,v in spec['command'].items());return r
 replayed=execute(fake,row['proposal'],row['checked'],row['fresh']);recorded=row['phases']
 if 'handoff' in replayed:replayed['handoff']['elapsed_s']=recorded['handoff']['elapsed_s']
 assert replayed==recorded and used==subset[3:]
from decision_pair_v1 import collect
turns=read(R/'turns.json')
for i,row in enumerate(turns):
 expected_source=read(R/'initial.json')['continuation']['observation'] if i==0 else turns[i-1]['decision_sampling']['result']['observation']
 assert row['source']==expected_source
 if 'decision_sampling' not in row:continue
 sampled=row['decision_sampling'];record=sampled['result'];checks=record['checks']
 subset=[c['result'] for c in calls[sampled['calls_begin']:sampled['calls_end']]]
 assert sampled['calls_begin']==row['calls_end'] and len(subset)==3*len(checks)
 observations=[]
 for j,check in enumerate(checks):
  triple=subset[3*j:3*j+3]
  assert [r['request']['command']['op'] for r in triple]==['clock','submit','clock']
  assert triple[1]['request']['command']['steps']==[{'op':'observe'}]
  assert triple[1]['state']['last_resolution']['terminal']['status']=='completed'
  assert check['fresh']==triple[1]['state']['continuation']['observation']
  assert check['clock']==triple[2]['state']['last_resolution']['clock']
  observations.append({'observation':check['fresh'],'image':str(T/Path(check['fresh']['image']).name),'clock':check['clock']})
 initial=calls[row['calls_end']-1]['result']['state']['continuation']['observation'];it=iter(observations)
 assert collect({'observation':initial,'image':str(T/Path(initial['image']).name)},lambda:next(it),3)==record
 if i+1<len(turns):
  feedback=turns[i+1]['feedback']
  assert feedback['resolution']==calls[row['calls_end']-1]['result']['state']['last_resolution']
  assert feedback['decision_samples']['status']==record['status']
terms=[e for e in events if e['event']=='terminal'];assert len(terms)==len([e for e in events if e['event']=='accepted'])
assert all(t['release']['verified'] and t['release']['keys_down']==[] and t['release']['buttons_down']==[] for t in terms)
assert all(e['admitted_ns']<e['valid_until_ns'] for e in events if e['event'] in ('input_admission','pointer_admission'))
decoder=Decoder('live-control');obs=[e for e in events if e['event']=='observation']
for n,e in enumerate(obs,1):
 f=decoder.accept((T/f'{n:03d}.ait').read_bytes())
 with Image.open(T/Path(e['image']).name) as im:assert (im.width,im.height,im.mode,im.tobytes())==(f.width,f.height,f.mode,f.pixels)
wb=load_workbook(T/'sheet.xlsx',data_only=True);values=[wb.active['A1'].value,wb.active['A2'].value];wb.close();assert values==[480,192]
evaluation=next(e for e in events if e['event']=='independent_evaluation');assert evaluation['success']
assert read(R/'result.json')['saved_success'] and read(R/'result.json')['exit_code']==0 and not (R/'error.json').exists()
assert not Path(ep['socket']).exists() and not Path(ep['cancel_socket']).exists()
snapshots=read(R/'workbook-snapshots.json')
expected_labels=['initial']+[label for r in turns if 'phases' in r for label in [f"after-phase-{r['turn']}",f"after-sampling-{r['turn']}"]]
assert [s['label'] for s in snapshots]==expected_labels
for snapshot in snapshots:
 assert snapshot['begin_ns']<=snapshot['finished_ns']
 if 'file' in snapshot:
  assert sha(R/snapshot['file'])==snapshot['sha256'] and (R/snapshot['file']).stat().st_size==snapshot['bytes']
 assert snapshot['read_status']=='bytes_copied_for_postrun_audit'
 wb=load_workbook(R/snapshot['file'],data_only=True);actual=[wb.active['A1'].value,wb.active['A2'].value];wb.close()
 snapshot['postrun_cells']=actual
 snapshot['matches_final_bytes']=snapshot['sha256']==sha(T/'sheet.xlsx')
 if snapshot['after_exchanges']:
  assert calls[snapshot['after_exchanges']-1]['end_ns']<=snapshot['begin_ns']
assert snapshots[0]['postrun_cells']==[None,None]
for i,row in enumerate(turns):
 if row['feedback'] is not None:
  assert set(row['feedback'])=={'proposal','phase_status','reason','target_check','resolution','decision_samples'}
 if 'phases' not in row:continue
 phase=next(s for s in snapshots if s['label']==f"after-phase-{row['turn']}")
 sampled=next(s for s in snapshots if s['label']==f"after-sampling-{row['turn']}")
 assert phase['after_exchanges']==row['calls_end'] and sampled['after_exchanges']==row['decision_sampling']['calls_end']
 if i+1<len(turns):assert sampled['finished_ns']<turns[i+1]['model_begin_ns']
report={'independent_saved_snapshots':snapshots,'decision_sampling':[{'turn':r['turn'],'elapsed_s':r['decision_sampling']['elapsed_s'],'captures':len(r['decision_sampling']['result']['checks']),'status':r['decision_sampling']['result']['status']} for r in turns if 'decision_sampling' in r], 'scope':'one actual Calc episode; shared caller; no cross-domain speed claim','events':len(events),'frames':len(obs),'exchanges':len(calls),'models':models,'saved_cells':values,'terminal_statuses':[t['status'] for t in terms],'capture_to_evaluation_s':(evaluation['known_ns']-obs[0]['capture_ns'])/1e9,'handoff_s':[r['phases']['handoff']['elapsed_s'] for r in read(R/'turns.json') if 'handoff' in r.get('phases',{})]}
(R/'audit.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8');print(json.dumps(report,indent=2))
