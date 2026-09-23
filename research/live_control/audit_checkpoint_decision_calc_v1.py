"""Audit actual model decisions, durable queries, and archived saved effects."""
import hashlib
import io
import json
import sys
from pathlib import Path

from PIL import Image
from openpyxl import load_workbook
from append_checkpoint_v1 import inspect, load
from calc_proposal_schema_v1 import parse
from checkpoint_contract_v1 import valid_reply
from decision_pair_v1 import collect
from durable_submit_v5 import reconcile
from effect_checkpoint import sample
from phased_submit_v1 import execute
from received_continuation_v1 import advance, start
from sampled_target_contract_v1 import evaluate

H = Path(__file__).resolve().parent
R = H / 'results/checkpoint-decision-calc-01'
T = R / 'runtime'
sys.path.insert(0, str(H.parent / 'observation_tiles'))
from tile_transport import Decoder


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


for name, digest in read(R / 'plan.json')['sources'].items():
    assert sha(H / name) == digest, name
for name, digest in read(T / 'sources.json').items():
    assert sha(H.parent / name) == digest, name
for name, digest in read(H / 'results/durable-checkpoint-controls-01/result.json')['sources'].items():
    assert sha(H / name) == digest, name

events = [json.loads(line) for line in (T / 'events.jsonl').read_bytes().splitlines()]
ep = read(R / 'endpoint.json')
state = start(ep['socket'])


def replay(result):
    global state
    q, reply = result['request'], result['reply']
    assert q['after'] == state['cursor']
    assert reply['records'] == events[q['after']:reply['cursor']]
    state = advance(state, ep['socket'], q['after'], reply)
    assert state == result.get('state', {}).get('continuation', result.get('continuation'))


initial = read(R / 'initial.json')
replay(initial)
calls = read(R / 'calls.json')
for call in calls:
    result = call['result']
    replay(result)
    q = result['request']
    pending = {'request': q, 'write_state': 'may_have_been_sent',
               'echo_seen': False, 'accepted': False, 'conflict': False}
    pending, resolution = reconcile(pending, result['reply']['records'])
    assert pending is None and result['state']['pending'] is None
    assert resolution == result['state']['last_resolution']
    assert result['state']['format'] == 'durable-submit-v5'
    assert result['reply']['command_receipt'] == {
        'request_id': q['request_id'], 'replayed': False, 'state': 'stdin_flushed'}
assert load(R / 'journal.jsonl') == calls[-1]['result']['state']
assert inspect(R / 'journal.jsonl')[1] == 1 + 2 * len(calls)
finish = read(R / 'finish.json')
replay(finish)
assert state['cursor'] == len(events)

turns = read(R / 'turns.json')
models = []
for i, row in enumerate(turns):
    n = row['turn']
    d = R / f'model-{n}'
    plan, process = read(d / 'plan.json'), read(d / 'process.json')
    source = initial['continuation']['observation'] if i == 0 else turns[i - 1]['decision_sampling']['result']['observation']
    assert row['source'] == source
    assert process['exit_code'] == 0
    assert plan['image_sha256'] == sha(T / Path(source['image']).name)
    assert plan['runner_sha256'] == sha(H / 'model_context_runner_v1.py')
    assert plan['instructions_sha256'] == sha(H / 'screenshot_responder_v1.txt')
    prompt = (d / 'prompt.txt').read_text(encoding='utf-8')
    assert prompt == (R / f'prompt-{n}.txt').read_text(encoding='utf-8')
    assert 'declared saved-cell checkpoint' in prompt
    assert json.loads(prompt.split('Evidence: ')[1]) == row['feedback']
    raw = (d / 'events.jsonl').read_bytes().splitlines(keepends=True)
    arrivals = [json.loads(line) for line in (d / 'arrivals.jsonl').read_bytes().splitlines()]
    assert len(raw) == len(arrivals) == 4
    for j, (line, arrival) in enumerate(zip(raw, arrivals)):
        assert arrival['line'] == j and arrival['bytes'] == len(line)
        assert arrival['sha256'] == hashlib.sha256(line).hexdigest()
    e = [json.loads(line) for line in raw]
    assert [item['type'] for item in e] == ['thread.started', 'turn.started', 'item.completed', 'turn.completed']
    assert e[2]['item']['type'] == 'agent_message'
    assert parse(e[2]['item']['text']) == row['proposal'] == read(R / f'proposal-{n}.json')
    models.append({'turn': n, 'usage': e[3]['usage'],
                   'runner_s': (process['exited_ns'] - process['started_ns']) / 1e9})
    assert row['proposal']['kind'] == 'act'
    subset = [c['result'] for c in calls[row['calls_begin']:row['calls_end']]]
    assert [r['request']['command']['op'] for r in subset[:3]] == ['clock', 'submit', 'clock']
    assert subset[1]['request']['command']['steps'] == [{'op': 'observe'}]
    assert row['fresh'] == subset[1]['state']['continuation']['observation']
    assert row['clock'] == subset[2]['state']['last_resolution']['clock']
    with Image.open(T / Path(source['image']).name) as old, Image.open(T / Path(row['fresh']['image']).name) as new:
        assert evaluate(row['contract'], {'intent': row['contract']['name'], 'execute_once': True},
                        source, row['fresh'], old, new, row['clock']['runtime_ns']) == row['checked']
    responses, used = iter(subset[3:]), []

    def fake(spec):
        result = next(responses)
        used.append(result)
        assert all(result['request']['command'][k] == v for k, v in spec['command'].items())
        return result

    phases = execute(fake, row['proposal'], row['checked'], row['fresh'])
    if 'handoff' in phases:
        phases['handoff']['elapsed_s'] = row['phases']['handoff']['elapsed_s']
    assert phases == row['phases'] and used == subset[3:]
    query = calls[row['calls_end']]['result']
    assert row['checkpoint'] == query
    assert query['state']['continuation']['observation'] == subset[-1]['state']['continuation']['observation']
    if 'decision_sampling' not in row:
        continue
    sampling = row['decision_sampling']
    record, checks = sampling['result'], sampling['result']['checks']
    assert sampling['calls_begin'] == row['calls_end'] + 1
    subset = [c['result'] for c in calls[sampling['calls_begin']:sampling['calls_end']]]
    assert len(subset) == 3 * len(checks)
    observations = []
    for j, check in enumerate(checks):
        triple = subset[3 * j:3 * j + 3]
        assert [r['request']['command']['op'] for r in triple] == ['clock', 'submit', 'clock']
        assert triple[1]['request']['command']['steps'] == [{'op': 'observe'}]
        assert triple[1]['state']['last_resolution']['terminal']['status'] == 'completed'
        assert check['fresh'] == triple[1]['state']['continuation']['observation']
        assert check['clock'] == triple[2]['state']['last_resolution']['clock']
        observations.append({'observation': check['fresh'], 'clock': check['clock'],
                             'image': str(T / Path(check['fresh']['image']).name)})
    before = calls[row['calls_end'] - 1]['result']['state']['continuation']['observation']
    samples = iter(observations)
    assert collect({'observation': before, 'image': str(T / Path(before['image']).name)}, lambda: next(samples), 3) == record
    feedback = turns[i + 1]['feedback']
    assert feedback['resolution'] == calls[row['calls_end'] - 1]['result']['state']['last_resolution']
    assert feedback['checkpoint_resolution'] == query['state']['last_resolution']
    assert feedback['decision_samples']['status'] == record['status']

policy = read(R / 'completion-policy.json')
contract = policy['contract']
assert policy['finish_on_verified'] is True
assert contract == {'kind': 'saved_cells', 'expected': {'A1': 480, 'A2': 192}}
queries = []
for i, call in enumerate(calls):
    result, q = call['result'], call['result']['request']
    if q['command']['op'] != 'effect_checkpoint':
        continue
    assert q['read_request_id'] == q['request_id']
    assert q['command']['contract'] == contract
    echo, event = result['reply']['records']
    assert echo['command'] == dict(q['command'], transport_request_id=q['request_id'])
    assert valid_reply(event, contract, q['request_id'])
    evidence = event['evidence']
    path = T / 'checkpoint-artifacts' / Path(evidence['archive_path']).name
    assert path.resolve() == Path(evidence['archive_path']).resolve()
    assert sha(path) == evidence['artifact_sha256'] and path.stat().st_size == evidence['archive_bytes']
    source = Path(evidence['archive_source_path'])
    assert source.name == 'sheet.xlsx' and source.parent.name.startswith('realapp-x-')
    assert not source.exists()
    reread = sample(path, contract)
    for field in ['status', 'actual', 'artifact_sha256', 'contract', 'reason', 'authority', 'task_success', 'observation_closed']:
        assert reread[field] == evidence[field], field
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False)
    actual = {address: wb.active[address].value for address in contract['expected']}
    wb.close()
    assert actual == evidence['actual']
    assert call['begin_ns'] <= echo['received_ns'] <= evidence['archive_started_ns'] <= evidence['source_sampled_ns']
    assert evidence['source_sampled_ns'] <= evidence['started_ns'] <= evidence['sampled_ns'] <= evidence['finished_ns'] <= event['emit_started_ns'] <= call['end_ns']
    queries.append({'call': i, 'status': evidence['status'], 'actual': actual,
                    'artifact_sha256': sha(path),
                    'roundtrip_ms': (call['end_ns'] - call['begin_ns']) / 1e6,
                    'received_to_emit_ms': (event['emit_started_ns'] - echo['received_ns']) / 1e6})
assert [q['status'] for q in queries] == ['UNKNOWN', 'VERIFIED']
assert queries[0]['actual'] == {'A1': None, 'A2': None}
assert queries[1]['actual'] == contract['expected']
assert turns[-1]['finished_by'] == 'verified_saved_contract'
assert len(turns) == 2 and queries[-1]['call'] == len(calls) - 1
assert 'decision_sampling' not in turns[-1]
assert finish['request']['command'] == {'op': 'finish'}
assert [e['event'] for e in finish['reply']['records']] == ['command', 'independent_evaluation']
unknown_cursor = calls[queries[0]['call']]['result']['reply']['cursor']
verified_cursor = calls[queries[1]['call']]['result']['reply']['cursor']
assert any(e['event'] == 'pointer_admission' for e in events[unknown_cursor:verified_cursor])
assert not any(e['event'] in ('input_admission', 'pointer_admission', 'accepted') for e in events[verified_cursor:])
commands = [c['result']['request']['command'] for c in calls]
save_steps = [s for c in commands for s in c.get('steps', []) if s == {'op': 'chord', 'modifier': 'Control_L', 'key': 's'}]
assert len(save_steps) == 1
terms = [e for e in events if e['event'] == 'terminal']
assert len(terms) == len([e for e in events if e['event'] == 'accepted'])
assert all(t['release']['verified'] and t['release']['keys_down'] == [] and t['release']['buttons_down'] == [] for t in terms)
assert all(e['admitted_ns'] < e['valid_until_ns'] for e in events if e['event'] in ('input_admission', 'pointer_admission'))
decoder = Decoder('live-control')
observations = [e for e in events if e['event'] == 'observation']
for n, event in enumerate(observations, 1):
    frame = decoder.accept((T / f'{n:03d}.ait').read_bytes())
    with Image.open(T / Path(event['image']).name) as image:
        assert (image.width, image.height, image.mode, image.tobytes()) == (frame.width, frame.height, frame.mode, frame.pixels)
wb = load_workbook(T / 'sheet.xlsx', data_only=False)
values = [wb.active['A1'].value, wb.active['A2'].value]
wb.close()
assert values == [480, 192]
evaluation = finish['reply']['records'][-1]
assert evaluation['success'] is True and evaluation['actual'] == values
assert evaluation['evaluation_request']['transport_request_id'] == finish['request']['request_id']
assert read(R / 'result.json') == {'exit_code': 0, 'model_calls': 2, 'saved_success': True,
                                  'completion_source': 'verified_saved_contract_then_independent_evaluation'}
assert not (R / 'error.json').exists()
assert not Path(ep['socket']).exists() and not Path(ep['cancel_socket']).exists()
report = {'scope': 'one actual Calc episode with an explicit artifact-evidence completion policy; not a pixel-only baseline',
          'events': len(events), 'frames': len(observations), 'exchanges': len(calls), 'models': models,
          'saved_cells': values, 'save_steps_submitted': len(save_steps), 'checkpoints': queries,
          'input_after_verified': 0, 'terminal_statuses': [t['status'] for t in terms],
          'capture_to_evaluation_s': (evaluation['known_ns'] - observations[0]['capture_ns']) / 1e9,
          'handoff_s': [r['phases']['handoff']['elapsed_s'] for r in turns if 'handoff' in r['phases']],
          'decision_sampling': [{'turn': r['turn'], 'elapsed_s': r['decision_sampling']['elapsed_s'],
                                 'captures': len(r['decision_sampling']['result']['checks'])}
                                for r in turns if 'decision_sampling' in r],
          'audit_source_sha256': sha(Path(__file__))}
(R / 'audit.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
print(json.dumps(report, indent=2))
