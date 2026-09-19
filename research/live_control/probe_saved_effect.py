"""Known retained artifacts and explicit unavailable/delayed/collateral controls."""
import hashlib,json,tempfile
from pathlib import Path
from openpyxl import Workbook
from saved_effect import inspect_saved_cells
HERE=Path(__file__).resolve().parent;out=HERE/'results/saved-effect-01';out.mkdir(exist_ok=False)
rows=[];sources={};contract={'A1':532,'A2':590}
for cohort in ('modal-target-route-01','modal-target-occlusion-01'):
    for row in json.loads((HERE/'results'/cohort/'results.json').read_text()):
        path=HERE/'results'/cohort/row['case']/'sheet.xlsx'
        evidence=inspect_saved_cells(path,action_id=row['case'],expected=contract,observation_closed=True)
        assert evidence['status']==('VERIFIED' if row['independent_evaluation']['success'] else 'CONTRADICTED')
        sources[str(path.relative_to(HERE))]=evidence['artifact_sha256'];rows.append(evidence)
with tempfile.TemporaryDirectory() as tmp:
    root=Path(tmp)
    for case in ('missing','corrupt','pending','formula','collateral','matching'):
        path=root/(case+'.xlsx');expected=dict(contract)
        if case=='corrupt':path.write_bytes(b'not an xlsx')
        elif case!='missing':
            wb=Workbook();ws=wb.active
            if case!='pending':ws['A1']=532;ws['A2']=590
            if case=='formula':ws['A1']='=532'
            if case=='collateral':ws['B1']='unwanted';expected['B1']=None
            wb.save(path);wb.close()
        evidence=inspect_saved_cells(path,action_id=case,expected=expected,observation_closed=case!='pending')
        assert evidence['status']==('CONTRADICTED' if case=='collateral' else 'VERIFIED' if case=='matching' else 'UNKNOWN')
        # Retain controls and stable error kind, not transient temporary paths.
        if path.exists():(out/path.name).write_bytes(path.read_bytes())
        if 'error' in evidence:evidence['error'].pop('message',None)
        rows.append(evidence)
for p in (Path(__file__),HERE/'saved_effect.py'):sources[str(p.relative_to(HERE))]=hashlib.sha256(p.read_bytes()).hexdigest()
(out/'sources.json').write_text(json.dumps(sources,indent=2)+'\n')
(out/'results.json').write_text(json.dumps(rows,indent=2)+'\n')
print(json.dumps([dict(case=r['action_id'],status=r['status'],reason=r['reason']) for r in rows],indent=2))
