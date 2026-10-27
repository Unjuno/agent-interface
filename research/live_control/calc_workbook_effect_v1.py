"""Post-controller Calc workbook scoring on a supplied immutable byte snapshot.

This helper never opens a GUI or selects an input action. The caller must obtain
and retain the snapshot and bind its digest to the current task independently.
"""
import hashlib
import io
import re
from openpyxl import load_workbook


def score_workbook(data, *, expected_digest, sheet_name, expected_values, evidence_ref):
    unavailable = {'status': 'unavailable'}
    if (type(data) is not bytes or not data or
            type(expected_digest) is not str or not re.fullmatch(r'[0-9a-f]{64}', expected_digest) or
            type(sheet_name) is not str or not sheet_name or
            type(expected_values) is not tuple or len(expected_values) != 2 or
            any(type(value) is not int for value in expected_values) or
            type(evidence_ref) is not str or not 0 < len(evidence_ref) <= 512):
        return unavailable
    digest = hashlib.sha256(data).hexdigest()
    if digest != expected_digest:
        return unavailable
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        if sheet_name not in workbook.sheetnames:
            return unavailable
        sheet = workbook[sheet_name]
        values = (sheet['A1'].value, sheet['A2'].value)
        matches = (all(type(value) is int for value in values) and values == expected_values)
    except Exception:
        return unavailable
    finally:
        if workbook is not None:
            workbook.close()
    return {'status': 'succeeded' if matches else 'failed',
            'evidence_ref': evidence_ref, 'evidence_digest': digest,
            'effect_scope': 'specified sheet A1 and A2 literal integer values'}