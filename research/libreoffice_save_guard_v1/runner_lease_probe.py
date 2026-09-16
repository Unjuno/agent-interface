#!/usr/bin/env python3
from pathlib import Path
import errno,fcntl,os,subprocess,time,json,hashlib,shutil
from openpyxl import Workbook
from Xlib import display
from office_backend import OfficeX11Backend
HERE=Path(__file__).resolve().parent
ROOT=HERE/'results'/'c284-lease-probe-01'
if ROOT.exists(): raise FileExistsError(ROOT)
ROOT.mkdir(parents=True)
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def statrow(p):
 s=Path(p).stat();return {'dev':s.st_dev,'ino':s.st_ino,'size':s.st_size,'mode':s.st_mode,'mtime_ns':s.st_mtime_ns,'ctime_ns':s.st_ctime_ns,'sha256':sha(p)}
def make_xlsx(p):
 wb=Workbook();ws=wb.active;ws['A1']='seed';ws['A2']='old';wb.save(p)
def wait_display(env,name):
 for _ in range(120):
  if subprocess.run(['xdpyinfo','-display',name],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0:return
  time.sleep(.05)
 raise RuntimeError('Xvfb not ready')
def windows(display_name):
 d=display.Display(display_name);root=d.screen().root;out=[]
 def walk(w):
  try:name=w.get_wm_name();cls=w.get_wm_class()
  except Exception:name=cls=None
  if name or cls:out.append({'id':w.id,'name':str(name) if name else None,'class':list(cls) if cls else None})
  try:children=w.query_tree().children
  except Exception:children=[]
  for c in children:walk(c)
 walk(root);d.close();return out
def find_calc(display_name):
 deadline=time.monotonic()+12
 while time.monotonic()<deadline:
  for r in windows(display_name):
   if r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']):return r['id']
  time.sleep(.05)
 raise RuntimeError('Calc window not found')
out=ROOT/'arm';out.mkdir();xlsx=out/'task.xlsx';make_xlsx(xlsx);auth=out/'Xauthority';auth.write_bytes(b'');disp=':286'
env=os.environ.copy();env['DISPLAY']=disp;env['XAUTHORITY']=str(auth);oldD=os.environ.get('DISPLAY');oldA=os.environ.get('XAUTHORITY');os.environ['DISPLAY']=disp;os.environ['XAUTHORITY']=str(auth)
xvfb=subprocess.Popen(['Xvfb',disp,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(out/'xvfb.stdout').open('w'),stderr=(out/'xvfb.stderr').open('w'));ob=lo=b=leasef=None
res={'allocation':'c284-lease-probe-01','display':disp,'initial_file':statrow(xlsx)}
try:
 wait_display(env,disp);ob=subprocess.Popen(['openbox'],env=env,stdout=(out/'openbox.stdout').open('w'),stderr=(out/'openbox.stderr').open('w'));profile=f'file://{out}/profile';lo=subprocess.Popen(['libreoffice',f'-env:UserInstallation={profile}','--nologo','--nodefault','--nofirststartwizard','--norestore',str(xlsx)],env=env,stdout=(out/'libreoffice.stdout').open('w'),stderr=(out/'libreoffice.stderr').open('w'))
 calc=find_calc(disp);res['calc_pid']=lo.pid;res['calc_window_id']=calc;time.sleep(.5);b=OfficeX11Backend(disp,{'calc':calc});b.focus('calc');b.pointer_move('calc','window_client',80,180);b.pointer_button('left',True);b.pointer_button('left',False);b.key_chord(['CTRL','Home']);b.text('office');b.key_chord(['ENTER']);b.text('preview');b.key_chord(['ENTER']);time.sleep(.15)
 d=display.Display(disp);focus=d.get_input_focus().focus;res['precheck']={'focus_id':getattr(focus,'id',None),'file':statrow(xlsx),'calc_alive':lo.poll() is None,'monotonic_ns':time.monotonic_ns()};d.close()
 leasef=xlsx.open('r+b',buffering=0);res['lease_fd_ino']=os.fstat(leasef.fileno()).st_ino
 try:
  rv=fcntl.fcntl(leasef.fileno(),fcntl.F_SETLEASE,fcntl.F_WRLCK);get=fcntl.fcntl(leasef.fileno(),fcntl.F_GETLEASE);res['lease']={'acquired':True,'set_return':rv,'get_lease':get,'errno':None,'error':None,'monotonic_ns':time.monotonic_ns()}
 except OSError as e:
  get=None
  try:get=fcntl.fcntl(leasef.fileno(),fcntl.F_GETLEASE)
  except OSError:pass
  res['lease']={'acquired':False,'set_return':None,'get_lease':get,'errno':e.errno,'error':e.strerror,'errno_name':errno.errorcode.get(e.errno),'monotonic_ns':time.monotonic_ns()}
 res['post_attempt']={'file':statrow(xlsx),'calc_alive':lo.poll() is None,'monotonic_ns':time.monotonic_ns()}
 rel=b.release_all().__dict__;res['release']=rel
 res['classification']='LEASE_ACQUIRED' if res['lease']['acquired'] else 'UNAVAILABLE_POST_PRECHECK'
 (ROOT/'RESULT.json').write_text(json.dumps(res,indent=2,default=str)+'\n');print(json.dumps(res,indent=2,default=str))
finally:
 if leasef:
  try:
   if res.get('lease',{}).get('acquired'):fcntl.fcntl(leasef.fileno(),fcntl.F_SETLEASE,fcntl.F_UNLCK)
  except Exception:pass
  leasef.close()
 if b:
  try:b.release_all();b.close()
  except Exception:pass
 if lo and lo.poll() is None:lo.terminate()
 if lo:
  try:lo.wait(timeout=3)
  except Exception:lo.kill()
 if ob and ob.poll() is None:ob.terminate()
 xvfb.terminate()
 if oldD is None:os.environ.pop('DISPLAY',None)
 else:os.environ['DISPLAY']=oldD
 if oldA is None:os.environ.pop('XAUTHORITY',None)
 else:os.environ['XAUTHORITY']=oldA
