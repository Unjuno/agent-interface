#!/usr/bin/env python3
import argparse,hashlib,json
from pathlib import Path
from openpyxl import load_workbook
ap=argparse.ArgumentParser();ap.add_argument('--xlsx',type=Path,required=True);ap.add_argument('--out',type=Path,required=True);a=ap.parse_args();p=a.xlsx
row={'exists':p.exists()}
if p.exists():
 b=p.read_bytes();row.update({'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b)})
 try:
  wb=load_workbook(p,data_only=False,read_only=True);ws=wb.active;row['cells']={k:ws[k].value for k in ('A1','A2','A3','B1')};row['xlsx_read_ok']=True
 except Exception as e:row['xlsx_read_ok']=False;row['error']=repr(e)
a.out.write_text(json.dumps(row,indent=2,default=str)+'\n');print(json.dumps(row,indent=2,default=str))
