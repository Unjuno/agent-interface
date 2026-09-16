#!/usr/bin/env python3
"""Post-outcome independent audit; does not import experiment runners/controllers."""
from __future__ import annotations
import copy, hashlib, json, sys
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
FLOCK=ROOT/'results'/'c284-flock-01'
LEASE=ROOT/'results'/'c284-lease-probe-01'
EXPECTED_BLOBS={'backend_x11.py':'b4f8e043ce4f8929d446e038418ea0fd3655bab0','office_backend.py':'3aeca10f62fb1366bcdd5fad561509cfcc0d91ab'}
def gitblob(p):
 b=p.read_bytes();return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def cells(p):
 wb=load_workbook(p,read_only=True,data_only=False);ws=wb.active;return {k:ws[k].value for k in ('A1','A2','A3','B1')}
def check(flock_summary,lease_result,*,filesystem=True):
 e=[]
 if flock_summary.get('allocation')!='c284-flock-01' or flock_summary.get('order')!=['stable','stale_restoremtime']:e.append('flock-summary-shape')
 rows={r['arm']:r for r in flock_summary.get('rows',[])}
 if set(rows)!={'stable','stale_restoremtime'}:e.append('flock-arm-set');return e
 st=rows['stable'];sr=rows['stale_restoremtime']
 for r in (st,sr):
  if not r.get('lock',{}).get('acquired'):e.append(r['arm']+':lock-not-acquired')
  if not r.get('release',{}).get('verified') or r['release'].get('keys_down') or r['release'].get('buttons_down'):e.append(r['arm']+':release')
  if r['after_ctrl_s']['modal_names']!=['Confirm File Format'] or r['after_optional_format_confirm']['modal_names']!=[]:e.append(r['arm']+':modal-sequence')
  if r['lock']['fd_ino']!=r['after_optional_format_confirm']['lock_fd_ino']:e.append(r['arm']+':lock-fd-moved')
  if r['after_optional_format_confirm']['file']['ino']==r['lock']['fd_ino']:e.append(r['arm']+':path-not-replaced')
 if st.get('classification')!='stable_saved':e.append('stable-class')
 if sr.get('classification')!='PERMISSIVE_STALE':e.append('stale-class')
 pre=sr['precheck']['file'];mut=sr['after_external_restoremtime']['file'];fin=sr['after_optional_format_confirm']['file']
 if mut['ino']!=pre['ino'] or mut['mtime_ns']!=pre['mtime_ns'] or mut['sha256']==pre['sha256']:e.append('stale-mutation-invariant')
 if fin['sha256']==mut['sha256']:e.append('stale-external-preserved-unexpected')
 if filesystem:
  for arm,r in rows.items():
   p=FLOCK/arm/'task.xlsx'
   if sha(p)!=r['after_optional_format_confirm']['file']['sha256']:e.append(arm+':durable-hash')
   c=cells(p)
   if c['A1']!='office' or c['A2']!='preview':e.append(arm+':durable-cells')
  repl=FLOCK/'stale_restoremtime'/'replacement.xlsx'
  rc=cells(repl)
  if rc!={'A1':'external','A2':'replacement','A3':'external-marker','B1':'writer'}:e.append('replacement-cells')
 if lease_result.get('allocation')!='c284-lease-probe-01':e.append('lease-allocation')
 if lease_result.get('classification')!='UNAVAILABLE_POST_PRECHECK':e.append('lease-class')
 le=lease_result.get('lease',{})
 if le.get('acquired') is not False or le.get('errno')!=11 or le.get('errno_name')!='EAGAIN' or le.get('get_lease')!=2:e.append('lease-eagain')
 if lease_result['precheck']['file']!=lease_result['post_attempt']['file']:e.append('lease-mutated-file')
 if not lease_result['precheck']['calc_alive'] or not lease_result['post_attempt']['calc_alive']:e.append('lease-calc-liveness')
 if not lease_result['release']['verified'] or lease_result['release']['keys_down'] or lease_result['release']['buttons_down']:e.append('lease-release')
 if filesystem:
  for n,x in EXPECTED_BLOBS.items():
   if gitblob(ROOT/n)!=x:e.append(n+':blob')
 return e
def main():
 fs=json.loads((FLOCK/'SUMMARY.json').read_text());lr=json.loads((LEASE/'RESULT.json').read_text())
 errors=check(fs,lr)
 corrupt=[]
 tests=[('stale-class',lambda a,b:a['rows'][1].__setitem__('classification','STALE_EFFECT_PREVENTED')),
        ('path-inode',lambda a,b:a['rows'][1]['after_optional_format_confirm']['file'].__setitem__('ino',a['rows'][1]['lock']['fd_ino'])),
        ('lease-errno',lambda a,b:b['lease'].__setitem__('errno',0)),
        ('release',lambda a,b:a['rows'][0]['release'].__setitem__('verified',False))]
 for name,mutator in tests:
  a,b=copy.deepcopy(fs),copy.deepcopy(lr);mutator(a,b);corrupt.append({'name':name,'rejected':bool(check(a,b,filesystem=False))})
 out={'pass':not errors and all(x['rejected'] for x in corrupt),'errors':errors,'flock_rows':2,'lease_rows':1,'corruptions':corrupt,
      'decision':{'flock':'PERMISSIVE_STALE','file_lease':'UNAVAILABLE_POST_PRECHECK'}}
 (ROOT/'AUDIT.json').write_text(json.dumps(out,indent=2)+'\n');print(json.dumps(out,indent=2));return 0 if out['pass'] else 1
if __name__=='__main__':raise SystemExit(main())
