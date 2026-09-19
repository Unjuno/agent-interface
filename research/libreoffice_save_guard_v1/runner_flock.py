#!/usr/bin/env python3
from pathlib import Path
import fcntl,os,subprocess,time,json,hashlib,shutil
from openpyxl import Workbook
from Xlib import display
from office_backend import OfficeX11Backend
HERE=Path(__file__).resolve().parent
ROOT=HERE/'results'/'c284-flock-01'
if ROOT.exists(): raise FileExistsError(ROOT)
ROOT.mkdir(parents=True)
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def statrow(p):
 s=Path(p).stat();return {'dev':s.st_dev,'ino':s.st_ino,'size':s.st_size,'mode':s.st_mode,'mtime_ns':s.st_mtime_ns,'ctime_ns':s.st_ctime_ns,'sha256':sha(p)}
def make_xlsx(p,a1,a2,a3=None,b1=None):
 wb=Workbook();ws=wb.active;ws['A1']=a1;ws['A2']=a2;ws['A3']=a3;ws['B1']=b1;wb.save(p)
def wait_display(env,name):
 for _ in range(120):
  if subprocess.run(['xdpyinfo','-display',name],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0:return
  time.sleep(.05)
 raise RuntimeError('Xvfb not ready')
def windows(display_name):
 d=display.Display(display_name);root=d.screen().root;out=[]
 def walk(w,depth=0):
  try:name=w.get_wm_name();cls=w.get_wm_class();geo=w.get_geometry();attrs=w.get_attributes()
  except Exception:name=cls=geo=attrs=None
  if name or cls:out.append({'id':w.id,'name':str(name) if name else None,'class':list(cls) if cls else None,'depth':depth,'w':getattr(geo,'width',None) if geo else None,'h':getattr(geo,'height',None) if geo else None,'map_state':getattr(attrs,'map_state',None) if attrs else None})
  try:children=w.query_tree().children
  except Exception:children=[]
  for c in children:walk(c,depth+1)
 walk(root);d.close();return out
def find_calc(display_name):
 deadline=time.monotonic()+12
 while time.monotonic()<deadline:
  for r in windows(display_name):
   if r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']):return r['id']
  time.sleep(.05)
 raise RuntimeError('Calc window not found')
def modal_names(rows,calc_id):
 return sorted(set(r['name'] for r in rows if r['id']!=calc_id and r['class'] and any('libreoffice-calc' in str(x).lower() for x in r['class']) and r['name']))
def run_arm(name,display_name,mutate):
 out=ROOT/name;out.mkdir();xlsx=out/'task.xlsx';make_xlsx(xlsx,'seed','old');auth=out/'Xauthority';auth.write_bytes(b'')
 env=os.environ.copy();env['DISPLAY']=display_name;env['XAUTHORITY']=str(auth);old_display=os.environ.get('DISPLAY');old_auth=os.environ.get('XAUTHORITY');os.environ['DISPLAY']=display_name;os.environ['XAUTHORITY']=str(auth)
 xvfb=subprocess.Popen(['Xvfb',display_name,'-screen','0','1024x768x24','-nolisten','tcp','-ac'],env=env,stdout=(out/'xvfb.stdout').open('w'),stderr=(out/'xvfb.stderr').open('w'))
 ob=lo=b=lockf=None;res={'arm':name,'mutate':mutate,'initial_file':statrow(xlsx)}
 try:
  wait_display(env,display_name);ob=subprocess.Popen(['openbox'],env=env,stdout=(out/'openbox.stdout').open('w'),stderr=(out/'openbox.stderr').open('w'))
  profile=f'file://{out}/profile';lo=subprocess.Popen(['libreoffice',f'-env:UserInstallation={profile}','--nologo','--nodefault','--nofirststartwizard','--norestore',str(xlsx)],env=env,stdout=(out/'libreoffice.stdout').open('w'),stderr=(out/'libreoffice.stderr').open('w'))
  calc=find_calc(display_name);res['calc_window_id']=calc;time.sleep(.5);b=OfficeX11Backend(display_name,{'calc':calc})
  b.focus('calc');b.pointer_move('calc','window_client',80,180);b.pointer_button('left',True);b.pointer_button('left',False);b.key_chord(['CTRL','Home']);b.text('office');b.key_chord(['ENTER']);b.text('preview');b.key_chord(['ENTER']);time.sleep(.15)
  d=display.Display(display_name);focus=d.get_input_focus().focus;focus_id=getattr(focus,'id',None);d.close();pre=statrow(xlsx);res['precheck']={'focus_id':focus_id,'calc_window_id':calc,'file':pre,'monotonic_ns':time.monotonic_ns()}
  if mutate:
   repl=out/'replacement.xlsx';make_xlsx(repl,'external','replacement','external-marker','writer');res['replacement_source']=statrow(repl);payload=repl.read_bytes()
   with xlsx.open('r+b') as f:f.seek(0);f.write(payload);f.truncate();f.flush();os.fsync(f.fileno())
   os.utime(xlsx,ns=(pre['mtime_ns'],pre['mtime_ns']))
   after=statrow(xlsx)
   if after['ino']!=pre['ino'] or after['mtime_ns']!=pre['mtime_ns'] or after['sha256']==pre['sha256']:raise RuntimeError('adversarial mutation invariant failed')
   res['after_external_restoremtime']={'file':after,'monotonic_ns':time.monotonic_ns()}
  lockf=xlsx.open('r+b',buffering=0);fcntl.flock(lockf.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB);lst=os.fstat(lockf.fileno())
  res['lock']={'acquired':True,'fd_ino':lst.st_ino,'path_before_save':statrow(xlsx),'monotonic_ns':time.monotonic_ns()}
  b.key_chord(['CTRL','S']);time.sleep(.7);first=windows(display_name);m1=modal_names(first,calc);res['after_ctrl_s']={'modal_names':m1,'file':statrow(xlsx),'lock_fd_ino':os.fstat(lockf.fileno()).st_ino,'monotonic_ns':time.monotonic_ns()}
  if m1==['Confirm File Format']:b.key_chord(['ENTER']);res['format_confirm_enter_sent']=True;time.sleep(1.3)
  else:res['format_confirm_enter_sent']=False
  second=windows(display_name);m2=modal_names(second,calc);res['after_optional_format_confirm']={'modal_names':m2,'file':statrow(xlsx),'lock_fd_ino':os.fstat(lockf.fileno()).st_ino,'monotonic_ns':time.monotonic_ns()}
  rel=b.release_all().__dict__;res['release']=rel;res['backend_emissions']=b.emissions
  scorep=out/'score.json';sp=subprocess.run(['python3',str(HERE/'score_any.py'),'--xlsx',str(xlsx),'--out',str(scorep)],capture_output=True,text=True);(out/'scorer.stdout').write_text(sp.stdout);(out/'scorer.stderr').write_text(sp.stderr);res['scorer_exitcode']=sp.returncode;res['score']=json.loads(scorep.read_text());cells=res['score'].get('cells') or {}
  if not mutate:res['classification']='stable_saved' if cells.get('A1')=='office' and cells.get('A2')=='preview' else 'stable_blocked_or_other'
  elif cells.get('A1')=='office' and cells.get('A2')=='preview':res['classification']='PERMISSIVE_STALE'
  elif cells.get('A1')=='external' and cells.get('A2')=='replacement':res['classification']='STALE_EFFECT_PREVENTED'
  else:res['classification']='UNCERTAIN_OTHER'
  (out/'result.json').write_text(json.dumps(res,indent=2,default=str)+'\n');return res
 finally:
  if lockf:
   try:fcntl.flock(lockf.fileno(),fcntl.LOCK_UN);lockf.close()
   except Exception:pass
  if b:
   try:b.release_all();b.close()
   except Exception:pass
  if lo and lo.poll() is None:lo.terminate()
  if lo:
   try:lo.wait(timeout=3)
   except Exception:lo.kill()
  if ob and ob.poll() is None:ob.terminate()
  xvfb.terminate()
  if old_display is None:os.environ.pop('DISPLAY',None)
  else:os.environ['DISPLAY']=old_display
  if old_auth is None:os.environ.pop('XAUTHORITY',None)
  else:os.environ['XAUTHORITY']=old_auth
rows=[run_arm('stable',':284','stable'=='stale'),run_arm('stale_restoremtime',':285',True)]
summary={'schema':'agent-interface/libreoffice-save-flock-v1','allocation':'c284-flock-01','order':['stable','stale_restoremtime'],'rows':rows,'source_blobs':{'backend_x11.py':'b4f8e043ce4f8929d446e038418ea0fd3655bab0','office_backend.py':'3aeca10f62fb1366bcdd5fad561509cfcc0d91ab'}}
(ROOT/'SUMMARY.json').write_text(json.dumps(summary,indent=2,default=str)+'\n');print(json.dumps({'allocation':summary['allocation'],'outcomes':[{'arm':r['arm'],'class':r['classification'],'m1':r['after_ctrl_s']['modal_names'],'m2':r['after_optional_format_confirm']['modal_names'],'lock_fd':r['lock']['fd_ino'],'path_ino_after':r['after_optional_format_confirm']['file']['ino'],'cells':r['score'].get('cells')} for r in rows]},indent=2))
