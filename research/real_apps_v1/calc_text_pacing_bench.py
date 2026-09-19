#!/usr/bin/env python3
import importlib.util,subprocess,time,json,random,statistics
from pathlib import Path
from openpyxl import Workbook,load_workbook
from Xlib import X,XK
from Xlib.ext import xtest
spec=importlib.util.spec_from_file_location('m',str(Path(__file__).with_name('real_app_suite_v1.py')));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
s=m.XSession();p=s.tmp/'sheet.xlsx';Workbook().save(p);D=s.d

def kc(n):return D.keysym_to_keycode(XK.string_to_keysym(n))
def raw(n,down):xtest.fake_input(D,X.KeyPress if down else X.KeyRelease,kc(n));D.sync()
def key(n):raw(n,1);raw(n,0)
def text(t,pause_ms):
 for i,ch in enumerate(t):
  raw(ch,1);raw(ch,0)
  if pause_ms and i+1<len(t):time.sleep(pause_ms/1000)
def chord(a,b):raw(a,1);raw(b,1);raw(b,0);raw(a,0)

try:
 s.spawn(['libreoffice','--norestore','--nodefault','--nolockcheck',f'-env:UserInstallation=file://{s.tmp}/lo','--calc',str(p)],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
 s.wait_window('sheet.xlsx',12);s.focus('sheet.xlsx');time.sleep(.35)
 rows=[]; expected=[]; rng=random.Random(913)
 pacings=[0,.1,.25,.5,1,2]
 for pacing in pacings:
  for j in range(12):
   # force duplicate runs while varying surrounding digits
   a=str(rng.randint(1,8));b=str(rng.randint(0,9));c=str(rng.randint(0,9))
   token=a+b+b+c+c+str(j%10)+str(j%10)
   expected.append((pacing,token));text(token,pacing);key('Return')
 chord('Control_L','s');m.wait_until(lambda:'Confirm File Format' in s.windows(),1,.002);time.sleep(.02);key('Return')
 time.sleep(.8)
 w=load_workbook(p,data_only=False,read_only=True);sh=w.active
 idx=1
 for pacing in pacings:
  vals=[];ok=0
  for j in range(12):
   exp=expected[idx-1][1]; got=str(sh.cell(idx,1).value); vals.append((exp,got));ok+=got==exp;idx+=1
  bad=[v for v in vals if v[0]!=v[1]]
  r={'char_gap_ms':pacing,'success':ok,'n':12,'rate':ok/12,'examples':bad[:3]};rows.append(r);print(json.dumps(r),flush=True)
 w.close(); out=Path(__file__).with_name('results'); out.mkdir(exist_ok=True); (out/'calc_text_pacing.json').write_text(json.dumps(rows,indent=2))
finally:s.close()
