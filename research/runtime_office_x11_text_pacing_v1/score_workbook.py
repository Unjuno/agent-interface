#!/usr/bin/env python3
from __future__ import annotations
import argparse, hashlib, json
from pathlib import Path
from openpyxl import load_workbook
CORPUS=['office','coffee','bookkeeper','committee','parallel','address','success','letterpress','mississippi','assessment','committee','bookkeeping','ffffffffff','aaaaaaaaaa','tttttttttt','ssssssssss']
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--xlsx',type=Path,required=True); ap.add_argument('--out',type=Path,required=True); a=ap.parse_args()
    wb=load_workbook(a.xlsx,data_only=False,read_only=True); ws=wb.active
    observed=[ws.cell(i,1).value for i in range(1,len(CORPUS)+1)]; next_value=ws.cell(len(CORPUS)+1,1).value
    mismatches=[{'row':i+1,'expected':e,'observed':o} for i,(e,o) in enumerate(zip(CORPUS,observed)) if e!=o]
    r={'schema':'agent-interface/office-x11-text-pacing-score-v1','passed':not mismatches and next_value is None,'expected':CORPUS,'observed':observed,'mismatches':mismatches,'next_value':next_value,'xlsx_sha256':hashlib.sha256(a.xlsx.read_bytes()).hexdigest(),'xlsx_bytes':a.xlsx.stat().st_size}
    a.out.write_text(json.dumps(r,indent=2)+'\n'); print(json.dumps(r,indent=2)); return 0 if r['passed'] else 1
if __name__=='__main__': raise SystemExit(main())
