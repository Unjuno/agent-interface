"""Construction-only readiness probe for the preregistered second desktop domain."""
import hashlib
import json
import os
from pathlib import Path
import signal
import subprocess
import tempfile
import time

from openpyxl import Workbook, load_workbook
from Xlib import display


root = Path(tempfile.mkdtemp(prefix="golden-v3-calc-"))
display_name = ":128"
xauth = root / "Xauthority"
xauth.touch(mode=0o600)
sheet = root / "fixture.xlsx"
wb = Workbook()
wb.active["A1"] = "construction"
wb.save(sheet)
env = os.environ.copy()
env.update(DISPLAY=display_name, XAUTHORITY=str(xauth))
xvfb = subprocess.Popen(["Xvfb", display_name, "-screen", "0", "1280x800x24", "-auth", str(xauth)], env=env)
calc = None
try:
    time.sleep(0.7)
    if xvfb.poll() is not None:
        raise RuntimeError("Xvfb exited")
    conn = display.Display(display_name)
    conn.sync()
    calc = subprocess.Popen([
        "libreoffice", "--norestore", "--nodefault", "--nolockcheck",
        f"-env:UserInstallation=file://{root / 'profile'}", "--calc", str(sheet)
    ], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    deadline = time.monotonic() + 12
    titles = []
    while time.monotonic() < deadline:
        conn.sync()
        titles = [w.get_wm_name() for w in conn.screen().root.query_tree().children if w.get_wm_name()]
        if any("fixture.xlsx" in title for title in titles):
            break
        time.sleep(0.1)
    else:
        raise RuntimeError(f"Calc window not observed; titles={titles}")
    checked = load_workbook(sheet, read_only=True, data_only=False).active["A1"].value
    print(json.dumps({
        "decision": "PASS_SECOND_DOMAIN_FIXTURE_READY",
        "fixture": "LibreOffice Calc",
        "display": display_name,
        "window_observed": True,
        "workbook_value": checked,
        "workbook_sha256": hashlib.sha256(sheet.read_bytes()).hexdigest(),
        "input_operations": 0,
        "formal_cases": 0,
    }, sort_keys=True))
finally:
    if calc is not None and calc.poll() is None:
        calc.send_signal(signal.SIGTERM)
        try:
            calc.wait(timeout=5)
        except subprocess.TimeoutExpired:
            calc.kill()
            calc.wait()
    if 'conn' in locals():
        conn.close()
    if xvfb.poll() is None:
        xvfb.terminate()
        xvfb.wait(timeout=5)
