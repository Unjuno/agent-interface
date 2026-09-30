"""One Calc construction row with same-run X-server window evidence."""
from __future__ import annotations

import hashlib
import json
import os
import pathlib
import re
import subprocess
import time

import uno
from openpyxl import Workbook, load_workbook

RUN = pathlib.Path("/out/calc-effect-contract-34-window-gate-20260927-01")
DISPLAY = ":97"
PORT = 2003
IMAGE = "issue-2849-task1-runtime@sha256:436172d89b145c6a9f9a57e655422c9558b3b0235347dd77607e9d61bcfa6393"


def sha256(path: pathlib.Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def connect_calc(deadline: float):
    local = uno.getComponentContext()
    resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
    while time.monotonic() < deadline:
        try:
            ctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={PORT};urp;StarOffice.ComponentContext")
            desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
            doc = desktop.getCurrentComponent()
            if doc is not None and doc.supportsService("com.sun.star.sheet.SpreadsheetDocument"):
                return doc
        except Exception:
            time.sleep(0.2)
    raise TimeoutError("Calc UNO document did not become available")


def command(argv: list[str]) -> str:
    p = subprocess.run(argv, text=True, capture_output=True, timeout=5)
    return json.dumps({"argv": argv, "exit": p.returncode, "stdout": p.stdout, "stderr": p.stderr}, sort_keys=True)


def main() -> int:
    if RUN.exists():
        raise FileExistsError(f"refusing to overwrite prior allocation: {RUN}")
    RUN.mkdir(parents=True)
    source = RUN / "baseline.xlsx"
    wb = Workbook()
    wb.active["A1"] = 0
    wb.active["B1"] = "baseline"
    wb.save(source)
    wb.close()
    before = sha256(source)

    os.environ["DISPLAY"] = DISPLAY
    xvfb = subprocess.Popen(["Xvfb", DISPLAY, "-screen", "0", "1280x800x24", "-nolisten", "tcp"], stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    time.sleep(0.5)
    profile = pathlib.Path("/tmp/calc-effect-contract-window-gate-profile")
    soffice = subprocess.Popen([
        "soffice", "--norestore", "--nofirststartwizard", "--nodefault", "--nologo",
        f"-env:UserInstallation={profile.as_uri()}",
        f"--accept=socket,host=127.0.0.1,port={PORT};urp;StarOffice.ServiceManager", str(source)
    ], stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, env=os.environ.copy())
    try:
        doc = connect_calc(time.monotonic() + 30)
        cell = doc.getSheets().getByIndex(0).getCellRangeByName("A1")
        before_live = cell.getValue()
        cell.setValue(7)
        live = cell.getValue()
        modified = bool(doc.isModified())
        tree = command(["xwininfo", "-root", "-tree"])
        tree_obj = json.loads(tree)
        match = re.search(r"(0x[0-9a-fA-F]+) \"VCL ImplGetDefaultWindow\"", tree_obj["stdout"])
        window_id = match.group(1) if match else None
        attrs = command(["xwininfo", "-id", window_id]) if window_id else None
        disk = load_workbook(source, read_only=True, data_only=True)
        persisted = disk.active["A1"].value
        disk.close()
        after = sha256(source)
        raw = {
            "schema": "calc-effect-contract-window-gate-raw-v1",
            "allocation": "calc-effect-contract-34-window-gate-20260927-01",
            "image": IMAGE,
            "application_version": subprocess.run(["soffice", "--version"], text=True, capture_output=True, timeout=10).stdout.strip(),
            "source_file": "baseline.xlsx",
            "source_sha256_before": before,
            "source_sha256_after": after,
            "cell_before_edit": before_live,
            "cell_after_edit_live": live,
            "document_modified_unsaved": modified,
            "save_store_calls": 0,
            "persisted_cell_after_independent_reopen": persisted,
            "xwininfo_root_tree": tree,
            "vcl_window_id": window_id,
            "xwininfo_window_attributes": attrs,
            "harness_completed": True,
            "scope": "one fresh construction case; no Agent Interface runtime or formal allocation",
        }
        (RUN / "raw.json").write_text(json.dumps(raw, sort_keys=True, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(raw, sort_keys=True))
        return 0
    finally:
        for proc in (soffice, xvfb):
            if proc.poll() is None:
                proc.kill()
                try:
                    proc.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    pass


if __name__ == "__main__":
    raise SystemExit(main())
